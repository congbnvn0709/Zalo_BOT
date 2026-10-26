"""
Unit and Integration Test Suite for FastAPI REST API & Web Dashboard (Milestone 3).
Uses starlette.testclient.TestClient for socket-free, fast, in-memory testing.
"""

import io
import os
import sys
import unittest
import openpyxl
from fastapi.testclient import TestClient

# Ensure project root in sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from main import app
from core.excel_integrator import DEFAULT_TEMPLATE_PATH


class TestApiIntegration(unittest.TestCase):
    """Integration test suite for FastAPI backend endpoints."""

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    def test_01_root_serves_html(self):
        """Test GET / serves HTML index with required DOM elements."""
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/html", response.headers["content-type"])
        content = response.text
        self.assertIn("<!DOCTYPE html>", content)
        self.assertIn('id="zalo-raw-text"', content)
        self.assertIn('id="btn-parse"', content)
        self.assertIn('id="btn-export"', content)

    def test_02_static_assets_served(self):
        """Test GET /static/css/app.css and /static/js/app.js are served properly."""
        css_res = self.client.get("/static/css/app.css")
        self.assertEqual(css_res.status_code, 200)
        self.assertIn("badge-valid", css_res.text)

        js_res = self.client.get("/static/js/app.js")
        self.assertEqual(js_res.status_code, 200)
        self.assertIn("handleParse", js_res.text)

    def test_03_api_health(self):
        """Test GET /api/health returns ok status."""
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data.get("status"), "ok")

    def test_04_api_sample(self):
        """Test GET /api/sample returns realistic sample text including HOÀNG ĐẮC PHÚ on 24/10."""
        response = self.client.get("/api/sample")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("sample_text", data)
        sample = data["sample_text"]
        self.assertIn("HOÀNG ĐẮC PHÚ", sample)
        self.assertIn("24/10", sample)

    def test_05_api_template_info(self):
        """Test GET /api/template/info returns template metadata."""
        response = self.client.get("/api/template/info")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("sheet_name", data)
        self.assertIn("employees", data)
        self.assertIn("date_range", data)
        self.assertIn("HOÀNG ĐẮC PHÚ", data["employees"])
        self.assertGreaterEqual(len(data["employees"]), 4)

    def test_06_api_parse_single_standard_message(self):
        """Test POST /api/parse correctly parses standard reference message."""
        msg = """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0"""
        response = self.client.post("/api/parse", json={"text": msg})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        records = data.get("records", [])
        summary = data.get("summary", {})

        self.assertEqual(len(records), 1)
        self.assertEqual(summary.get("total"), 1)
        self.assertEqual(summary.get("valid") + summary.get("warning"), 1)
        self.assertEqual(summary.get("error"), 0)

        rec = records[0]
        self.assertEqual(rec["employee_name"], "HOÀNG ĐẮC PHÚ")
        self.assertEqual(rec["hours"]["hcn_100"], 8.0)
        self.assertEqual(rec["hours"]["tcn_150"], 0.0)
        self.assertEqual(rec["meal_count"], 0)
        self.assertIn(rec["status"], ("VALID", "WARNING"))

    def test_06b_api_parse_full_date_valid(self):
        """Test POST /api/parse with explicit full date produces VALID status."""
        msg = """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10/2026
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0"""
        response = self.client.post("/api/parse", json={"text": msg})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        summary = data.get("summary", {})
        self.assertEqual(summary.get("valid"), 1)
        self.assertEqual(data["records"][0]["status"], "VALID")

    def test_07_api_parse_batch_mixed_messages(self):
        """Test POST /api/parse handles multi-employee batch with shifts and overtime."""
        sample_res = self.client.get("/api/sample")
        sample_text = sample_res.json()["sample_text"]

        response = self.client.post("/api/parse", json={"text": sample_text})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        records = data.get("records", [])
        summary = data.get("summary", {})

        self.assertGreaterEqual(len(records), 5)
        self.assertEqual(summary["total"], len(records))
        self.assertGreaterEqual(summary["valid"] + summary["warning"], 5)

        # Verify overtime day record (NGUYỄN NGỌC THAO: 8h-20h -> 8h HCN, 3h TCN)
        thao_recs = [r for r in records if "THAO" in r["employee_name"]]
        self.assertTrue(len(thao_recs) > 0)
        thao = thao_recs[0]
        self.assertEqual(thao["hours"]["hcn_100"], 8.0)
        self.assertEqual(thao["hours"]["tcn_150"], 3.0)
        self.assertEqual(thao["meal_count"], 1)

    def test_08_api_parse_empty_payload(self):
        """Test POST /api/parse with empty text returns total=0 and empty records list."""
        response = self.client.post("/api/parse", json={"text": ""})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["summary"]["total"], 0)
        self.assertEqual(data["records"], [])

    def test_09_api_parse_corrupt_payload(self):
        """Test POST /api/parse with unstructured noise returns status ERROR without crashing."""
        noisy_text = "Hôm nay tôi mệt quá xin nghỉ ốm một hôm nhé sếp"
        response = self.client.post("/api/parse", json={"text": noisy_text})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["summary"]["total"], 1)
        self.assertEqual(data["summary"]["error"], 1)
        self.assertEqual(data["records"][0]["status"], "ERROR")

    def test_10_api_export_excel_success_and_download(self):
        """Test full pipeline: parse -> export -> download token -> download stream -> check cells & formulas."""
        msg = """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0"""
        parse_res = self.client.post("/api/parse", json={"text": msg})
        records = parse_res.json()["records"]

        export_res = self.client.post(
            "/api/export-excel",
            json={"records": records, "output_filename": "TEST_HOANG_DAC_PHU.xlsx"}
        )
        self.assertEqual(export_res.status_code, 200)
        export_data = export_res.json()
        self.assertTrue(export_data["success"])
        self.assertIn("download_token", export_data)
        token = export_data["download_token"]
        self.assertIn("summary", export_data)
        self.assertEqual(export_data["summary"]["cells_written"], 2)

        # Download using token
        dl_res = self.client.get(f"/api/download/{token}")
        self.assertEqual(dl_res.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            dl_res.headers["content-type"],
        )
        self.assertIn('attachment; filename="TEST_HOANG_DAC_PHU.xlsx"', dl_res.headers["content-disposition"])

        # Open downloaded bytes with openpyxl and verify cells
        file_bytes = io.BytesIO(dl_res.content)
        wb = openpyxl.load_workbook(file_bytes, data_only=False)
        ws = wb.active

        # Row 24 (HCN 100% of HOÀNG ĐẮC PHÚ), Column 11 (2026-10-24)
        self.assertEqual(ws.cell(row=24, column=11).value, 8)
        # Row 30 (Suất ăn of HOÀNG ĐẮC PHÚ), Column 11
        self.assertEqual(ws.cell(row=30, column=11).value, 0)
        # Formula preservation in row 24 column AK (37)
        self.assertEqual(ws.cell(row=24, column=37).value, "=SUM(F24:AJ24)")
        # Full calculation flag
        self.assertTrue(wb.calculation.fullCalcOnLoad)

    def test_11_api_download_invalid_token(self):
        """Test GET /api/download with nonexistent token returns 404."""
        response = self.client.get("/api/download/nonexistent_fake_token_123")
        self.assertEqual(response.status_code, 404)

    def test_12_api_export_excel_empty_records(self):
        """Test POST /api/export-excel with empty records raises 400."""
        response = self.client.post("/api/export-excel", json={"records": []})
        self.assertEqual(response.status_code, 400)

    def test_13_api_template_upload_wrong_extension(self):
        """Test POST /api/template/upload with non-xlsx file returns 400."""
        file_content = b"fake csv content"
        response = self.client.post(
            "/api/template/upload",
            files={"file": ("test.txt", file_content, "text/plain")}
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("xlsx", response.json()["detail"].lower())

    def test_14_api_template_upload_valid_xlsx(self):
        """Test POST /api/template/upload with valid xlsx returns updated template metadata."""
        # Read default template as upload source
        with open(DEFAULT_TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()

        response = self.client.post(
            "/api/template/upload",
            files={"file": ("custom_sample.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("sheet_name", data)
        self.assertIn("employees", data)
        self.assertIn("HOÀNG ĐẮC PHÚ", data["employees"])

    def test_15_api_export_with_custom_template_flag(self):
        """Test POST /api/export-excel with use_custom_template=True works correctly."""
        msg = """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0"""
        parse_res = self.client.post("/api/parse", json={"text": msg})
        records = parse_res.json()["records"]

        export_res = self.client.post(
            "/api/export-excel",
            json={"records": records, "use_custom_template": True}
        )
        self.assertEqual(export_res.status_code, 200)
        export_data = export_res.json()
        self.assertTrue(export_data["success"])
        self.assertIn("download_token", export_data)


if __name__ == "__main__":
    unittest.main()
