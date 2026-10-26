"""
End-to-End (E2E) Test Suite for Zalo Attendance Web App & Excel Integrator.

Covers:
- Tier 3: Cross-Feature Combinations (Pairwise & Multi-Feature Interactions >= 11 tests)
- Tier 4: Real-World Application Scenarios (>= 5 tests, including exact Acceptance Criteria test case)

Total: >= 16 test cases.
Follows opaque-box, requirement-driven methodology derived from ORIGINAL_REQUEST.md,
PROJECT.md, and TEST_INFRA.md.
"""

import os
import io
import re
import csv
import copy
import unicodedata
from difflib import SequenceMatcher
from typing import List, Dict, Any, Optional

import pytest
import openpyxl

from models.attendance import WorkHours, AttendanceRecord, ParseSummary, ParseResponse
from core.parser import (
    AttendanceParser,
    parse_attendance_messages,
    parse_single_message,
    normalize_vietnamese_name,
    calculate_work_hours,
)


PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SAMPLE_CSV_PATH = os.path.join(PROJECT_ROOT, "sample_data.csv")


# ==============================================================================
# E2E HARNESS & WORKBOOK ENGINE HELPERS
# ==============================================================================

def strip_accents(text: str) -> str:
    if not text:
        return ""
    text = text.replace('đ', 'd').replace('Đ', 'D')
    nfkd = unicodedata.normalize('NFKD', text)
    return ''.join(c for c in nfkd if not unicodedata.combining(c))


def resolve_employee_row(name: str) -> Optional[int]:
    registry = {
        "NGUYỄN NGỌC THAO": 3,
        "NGUYỄN ĐỨC CÔNG": 10,
        "TÔ VĂN HƯNG": 17,
        "HOÀNG ĐẮC PHÚ": 24,
    }
    clean = re.sub(r'\s+', ' ', name.strip()).upper()
    for reg, r in registry.items():
        if clean == reg:
            return r
    clean_no_acc = strip_accents(clean)
    for reg, r in registry.items():
        if clean_no_acc == strip_accents(reg):
            return r
    for reg, r in registry.items():
        if SequenceMatcher(None, clean_no_acc, strip_accents(reg)).ratio() >= 0.85:
            return r
    return None


def resolve_date_column(date_str: str) -> Optional[int]:
    with open(SAMPLE_CSV_PATH, 'r', encoding='utf-8') as f:
        header = next(csv.reader(f))
    m_iso = re.search(r'(\d{4})[-/](\d{1,2})[-/](\d{1,2})', date_str)
    if m_iso:
        t_day, t_mon = int(m_iso.group(3)), int(m_iso.group(2))
    else:
        m = re.search(r'(\d{1,2})[-/](\d{1,2})', date_str)
        if m:
            t_day, t_mon = int(m.group(1)), int(m.group(2))
        else:
            return None

    for col_idx, cell in enumerate(header, start=1):
        if col_idx < 6:
            continue
        m_col = re.match(r'(\d{4})-(\d{1,2})-(\d{1,2})', cell)
        if m_col:
            if int(m_col.group(3)) == t_day and int(m_col.group(2)) == t_mon:
                return col_idx
    return None


def get_fresh_workbook() -> openpyxl.Workbook:
    """Builds a pristine in-memory workbook loaded from sample_data.csv with live formulas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CHẤM CÔNG WEVN - ETS"
    wb.calculation.fullCalcOnLoad = True

    with open(SAMPLE_CSV_PATH, 'r', encoding='utf-8') as f:
        for r_idx, row in enumerate(csv.reader(f), start=1):
            for c_idx, val in enumerate(row, start=1):
                if val != "":
                    try:
                        n = float(val)
                        ws.cell(row=r_idx, column=c_idx, value=int(n) if n.is_integer() else n)
                    except ValueError:
                        ws.cell(row=r_idx, column=c_idx, value=val)

    # Individual formulas
    for r in [3, 10, 17, 24]:
        ws.cell(row=r, column=37, value=f"=SUM(F{r}:AJ{r})")
        ws.cell(row=r, column=38, value=f"=SUM(F{r+1}:AJ{r+1})")
        ws.cell(row=r, column=39, value=f"=SUM(F{r+2}:AJ{r+2})")
        ws.cell(row=r, column=40, value=f"=SUM(F{r+3}:AJ{r+3})")
        ws.cell(row=r, column=41, value=f"=SUM(F{r+4}:AJ{r+4})")
        ws.cell(row=r, column=42, value=f"=SUM(F{r+5}:AJ{r+5})")
        ws.cell(row=r, column=43, value=f"=SUM(F{r+6}:AJ{r+6})")
        ws.cell(row=r, column=44, value=f"=SUM(AK{r}:AP{r})")

    # Grand totals
    for col in range(37, 45):
        let = openpyxl.utils.get_column_letter(col)
        ws.cell(row=2, column=col, value=f"=SUM({let}3,{let}10,{let}17,{let}24)")

    return wb


def apply_attendance_records_to_workbook(wb: openpyxl.Workbook, records: List[AttendanceRecord]) -> Dict[str, Any]:
    """
    Applies AttendanceRecords to openpyxl workbook preserving all formulas.
    Returns audit statistics.
    """
    ws = wb.active
    updated_cells = 0
    employees_updated = set()
    errors = []

    for rec in records:
        if rec.status == "ERROR":
            errors.append(f"Skipping record with error: {rec.messages}")
            continue

        r_start = resolve_employee_row(rec.normalized_name or rec.employee_name)
        if not r_start:
            errors.append(f"Unmatched employee: {rec.employee_name}")
            continue

        col_date = resolve_date_column(rec.date_str or rec.resolved_date)
        if not col_date:
            errors.append(f"Date out of range: {rec.date_str}")
            continue

        # Map work hours to row offsets
        # Offset 0: HCN 100%
        # Offset 1: TCN 150%
        # Offset 2: HCĐ 130%
        # Offset 3: TCĐ 195%
        # Offset 4: Ngày CN (200%)
        # Offset 5: Đêm CN (260%)
        # Offset 6: Suất ăn
        h = rec.hours
        if h.hcn_100 > 0:
            ws.cell(row=r_start + 0, column=col_date, value=int(h.hcn_100) if h.hcn_100.is_integer() else h.hcn_100)
            updated_cells += 1
        if h.tcn_150 > 0:
            ws.cell(row=r_start + 1, column=col_date, value=int(h.tcn_150) if h.tcn_150.is_integer() else h.tcn_150)
            updated_cells += 1
        if h.hcd_130 > 0:
            ws.cell(row=r_start + 2, column=col_date, value=int(h.hcd_130) if h.hcd_130.is_integer() else h.hcd_130)
            updated_cells += 1
        if h.tcd_195 > 0:
            ws.cell(row=r_start + 3, column=col_date, value=int(h.tcd_195) if h.tcd_195.is_integer() else h.tcd_195)
            updated_cells += 1
        if h.sunday_day_200 > 0:
            ws.cell(row=r_start + 4, column=col_date, value=int(h.sunday_day_200) if h.sunday_day_200.is_integer() else h.sunday_day_200)
            updated_cells += 1
        if h.sunday_night_260 > 0:
            ws.cell(row=r_start + 5, column=col_date, value=int(h.sunday_night_260) if h.sunday_night_260.is_integer() else h.sunday_night_260)
            updated_cells += 1

        # Meal count
        ws.cell(row=r_start + 6, column=col_date, value=rec.meal_count)
        updated_cells += 1

        employees_updated.add(r_start)

    return {
        "updated_cells": updated_cells,
        "employees_count": len(employees_updated),
        "errors": errors
    }


# ==============================================================================
# TIER 3: CROSS-FEATURE COMBINATIONS (>= 11 test cases)
# ==============================================================================

import unittest


class TestTier3CrossFeatureCombinations(unittest.TestCase):
    """Tier 3: Pairwise & Cross-Feature Interaction Tests."""

    def test_tier3_01_night_shift_and_overtime(self):
        """Cross-feature: F01 (Name) + F03 (Night Overtime: HCĐ 130% + TCĐ 195%)."""
        msg = """HỌ TÊN: NGUYỄN ĐỨC CÔNG
NGÀY: 24/10
CA: Ca đêm
DỰ ÁN: ETS
THỜI GIAN LÀM VIỆC: 20h-8h (11H)
CÓ ĂN Ở XƯỞNG KO: 1"""
        rec = parse_single_message(msg)
        assert rec.normalized_name.upper() == "NGUYỄN ĐỨC CÔNG"
        assert rec.hours.hcd_130 == 8.0
        assert rec.hours.tcd_195 == 3.0
        assert rec.hours.hcn_100 == 0.0
        assert rec.hours.tcn_150 == 0.0
        assert rec.hours.total_hours == 11.0
        assert rec.meal_count == 1

    def test_tier3_02_sunday_day_and_meal(self):
        """Cross-feature: F02 (Sunday Resolution) + F03 (Sunday Rate 200%) + F04 (Meal)."""
        msg = """HỌ TÊN: HOÀNG ĐẮC PHÚ
NGÀY: 25/10/2026
CA: Ca ngày
THỜI GIAN LÀM VIỆC: 8h-17h (8H)
CÓ ĂN Ở XƯỞNG KO: Có"""
        rec = parse_single_message(msg)
        assert rec.is_sunday is True
        assert rec.hours.sunday_day_200 == 8.0
        assert rec.hours.hcn_100 == 0.0
        assert rec.meal_count == 1

    def test_tier3_03_sunday_night_and_overtime(self):
        """Cross-feature: F02 (Sunday) + F03 (Sunday Night 260%) + F07 (Coordinate Row Offset 5)."""
        msg = """HỌ TÊN: HOÀNG ĐẮC PHÚ
NGÀY: 25/10/2026
CA: Ca đêm
THỜI GIAN LÀM VIỆC: 20h-8h (11H)
CÓ ĂN Ở XƯỞNG KO: 1"""
        rec = parse_single_message(msg)
        assert rec.is_sunday is True
        assert rec.hours.sunday_night_260 == 11.0
        row = resolve_employee_row(rec.normalized_name)
        assert row == 24
        # Sunday night row offset is 5 -> row 29
        target_row = row + 5
        assert target_row == 29

    def test_tier3_04_batch_with_error_and_warning_tolerance(self):
        """Cross-feature: F05 (Batch Parsing) + F09 (Status Classification: VALID, WARNING, ERROR)."""
        batch = """HỌ TÊN: HOÀNG ĐẮC PHÚ
NGÀY: 24/10/2026
CA: Ca ngày
THỜI GIAN: 8h-17h (8H)
CÓ ĂN: 0

HỌ TÊN: NGUYỄN NGỌC THAO
NGÀY: 24/10
CA: Ca ngày
THỜI GIAN: 8h-20h (11H)
CÓ ĂN: 1

HỌ TÊN: TÔ VĂN HƯNG
CA: Ca đêm
THỜI GIAN: 20h-5h (8H)
CÓ ĂN: 0"""
        records = parse_attendance_messages(batch)
        assert len(records) == 3
        # Record 1: VALID (explicit year)
        assert records[0].status == "VALID"
        # Record 2: WARNING (inferred year 2026)
        assert records[1].status == "WARNING"
        # Record 3: ERROR (missing date)
        assert records[2].status == "ERROR"

    def test_tier3_05_desktop_timestamp_and_fuzzy_name(self):
        """Cross-feature: F01 (Fuzzy/Unaccented Match) + F05 (Metadata Strip) + F07 (Row Resolver)."""
        raw = """[17:35, 24/10/2026] Phu Hoang: HỌ TÊN : HOANG DAC PHU
NGÀY: 24/10
CA: Ca ngày
THỜI GIAN: 8h-17h (8H)
CÓ ĂN: 0"""
        rec = parse_single_message(raw)
        assert "HOANG DAC PHU" in rec.normalized_name.upper()
        row = resolve_employee_row(rec.normalized_name)
        assert row == 24

    def test_tier3_06_single_line_hyphen_and_excel_write(self):
        """Cross-feature: F05 (Single-Line Parse) + F08 (Excel Numeric Write & Formula Preservation)."""
        msg = "HỌ TÊN: NGUYỄN NGỌC THAO - NGÀY: 24/10 - CA: Ca ngày - THỜI GIAN: 8h-17h (8H) - ĂN: 0"
        rec = parse_single_message(msg)
        assert rec.hours.hcn_100 == 8.0

        wb = get_fresh_workbook()
        stats = apply_attendance_records_to_workbook(wb, [rec])
        assert stats["updated_cells"] >= 1
        ws = wb.active
        # NGUYỄN NGỌC THAO row 3, col 11 (24/10)
        assert ws.cell(row=3, column=11).value == 8
        assert ws.cell(row=3, column=37).value == "=SUM(F3:AJ3)"

    def test_tier3_07_unaccented_input_and_formula_recalculation(self):
        """Cross-feature: F01 (Unaccented) + F02 (Date) + F08 (Excel Integration)."""
        msg = """HO TEN : HOANG DAC PHU
NGAY: 24/10
CA: Ca ngay
THOI GIAN: 8h-17h (8H)
AN: 0"""
        rec = parse_single_message(msg)
        wb = get_fresh_workbook()
        apply_attendance_records_to_workbook(wb, [rec])
        ws = wb.active
        # Cell K24 = 8, Formula in AK24 intact
        assert ws.cell(row=24, column=11).value == 8
        assert ws.cell(row=24, column=37).value == "=SUM(F24:AJ24)"

    def test_tier3_08_inferred_shift_from_time_window(self):
        """Cross-feature: F03 (Shift Inference from 20h start) + F07 (Offset 2 HCĐ)."""
        msg = """HỌ TÊN: NGUYỄN ĐỨC CÔNG
NGÀY: 24/10/2026
THỜI GIAN LÀM VIỆC: 20h-5h (8H)
CÓ ĂN Ở XƯỞNG KO: 0"""
        rec = parse_single_message(msg)
        assert rec.shift_type == "NIGHT"
        assert rec.hours.hcd_130 == 8.0
        assert rec.hours.hcn_100 == 0.0

    def test_tier3_09_custom_template_and_coordinate_export(self):
        """Cross-feature: F06 (Template Generation) + F08 (Cell Writing)."""
        wb = get_fresh_workbook()
        rec = parse_single_message("HỌ TÊN: TÔ VĂN HƯNG\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nCÓ ĂN: 0")
        apply_attendance_records_to_workbook(wb, [rec])
        ws = wb.active
        # TÔ VĂN HƯNG is row 17, col 11
        assert ws.cell(row=17, column=11).value == 8

    def test_tier3_10_discrepant_bracket_and_warning(self):
        """Cross-feature: F03 (Bracketed Hour Precedence) + F05 (Warning Logging)."""
        msg = """HỌ TÊN: HOÀNG ĐẮC PHÚ
NGÀY: 24/10/2026
CA: Ca ngày
THỜI GIAN LÀM VIỆC: 8h-17h (10H)
CÓ ĂN: 0"""
        rec = parse_single_message(msg)
        assert rec.hours.total_hours == 10.0
        assert any("lớn hơn" in m.lower() for m in rec.messages)
        assert rec.status == "WARNING"

    def test_tier3_11_double_meal_and_long_overtime(self):
        """Cross-feature: F03 (Overtime) + F04 (Double Meal=2) + F08 (Write Offset 6)."""
        msg = """HỌ TÊN: HOÀNG ĐẮC PHÚ
NGÀY: 24/10/2026
CA: Ca ngày
THỜI GIAN LÀM VIỆC: 8h-20h (11H)
CÓ ĂN Ở XƯỞNG KO: 2"""
        rec = parse_single_message(msg)
        assert rec.meal_count == 2
        wb = get_fresh_workbook()
        apply_attendance_records_to_workbook(wb, [rec])
        ws = wb.active
        # Meal is row 24 + 6 = 30, col 11
        assert ws.cell(row=30, column=11).value == 2


# ==============================================================================
# TIER 4: REAL-WORLD APPLICATION SCENARIOS (>= 5 test cases)
# ==============================================================================

class TestTier4RealWorldScenarios(unittest.TestCase):
    """Tier 4: Realistic Workload & End-to-End Application Scenarios."""

    def test_tier4_01_acceptance_criteria_hoang_dac_phu(self):
        """
        Scenario 1: Official Acceptance Criteria Verification.
        Exact message from ORIGINAL_REQUEST.md:
        HỌ TÊN : HOÀNG ĐẮC PHÚ
        NGÀY: 24/10
        CA : Ca ngày 
        DỰ ÁN : sdv 
        THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
        CÓ ĂN Ở XƯỞNG KO : 0
        """
        raw_msg = """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0"""

        # 1. Parse
        rec = parse_single_message(raw_msg)
        assert rec.normalized_name.upper() == "HOÀNG ĐẮC PHÚ"
        assert rec.day_number == 24
        assert rec.shift_type == "DAY"
        assert rec.hours.hcn_100 == 8.0
        assert rec.hours.tcn_150 == 0.0
        assert rec.hours.hcd_130 == 0.0
        assert rec.hours.tcd_195 == 0.0
        assert rec.meal_count == 0

        # 2. Excel Integration
        wb = get_fresh_workbook()
        ws = wb.active

        # Verify prior state: In sample_data.csv, HOÀNG ĐẮC PHÚ has 8 on 2026-10-22 (Col I, index 9)
        assert ws.cell(row=24, column=9).value == 8
        # Cell K24 (col 11) is initially None or empty
        assert ws.cell(row=24, column=11).value is None or ws.cell(row=24, column=11).value == ""

        # Apply record
        audit = apply_attendance_records_to_workbook(wb, [rec])
        assert audit["employees_count"] == 1
        assert len(audit["errors"]) == 0

        # 3. Assert cell K24 = 8 (numeric) and K30 = 0 (numeric)
        assert ws.cell(row=24, column=11).value == 8
        assert ws.cell(row=24, column=11).data_type == 'n'
        assert ws.cell(row=30, column=11).value == 0
        assert ws.cell(row=30, column=11).data_type == 'n'

        # 4. Assert summary formula in AK24 and AR24 remain formulas
        assert ws.cell(row=24, column=37).value == "=SUM(F24:AJ24)"
        assert ws.cell(row=24, column=44).value == "=SUM(AK24:AP24)"

        # 5. Assert workbook configuration
        assert wb.calculation.fullCalcOnLoad is True

        # 6. Verify save and reload with data_only=False
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        reloaded_wb = openpyxl.load_workbook(buf, data_only=False)
        r_ws = reloaded_wb.active
        assert r_ws.cell(row=24, column=11).value == 8
        assert r_ws.cell(row=30, column=11).value == 0
        assert r_ws.cell(row=24, column=37).value == "=SUM(F24:AJ24)"

    def test_tier4_02_multi_employee_mixed_shifts(self):
        """
        Scenario 2: Multi-Employee Factory Shift Roster.
        Batch with all 4 employees covering Day, Overtime, Night, and Sunday shifts.
        """
        roster_text = """HỌ TÊN: NGUYỄN NGỌC THAO
NGÀY: 24/10/2026
CA: Ca ngày
THỜI GIAN: 8h-20h (11H)
CÓ ĂN: 1

HỌ TÊN: NGUYỄN ĐỨC CÔNG
NGÀY: 24/10/2026
CA: Ca đêm
THỜI GIAN: 20h-5h (8H)
CÓ ĂN: 0

HỌ TÊN: TÔ VĂN HƯNG
NGÀY: 24/10/2026
CA: Ca đêm
THỜI GIAN: 20h-8h (11H)
CÓ ĂN: 1

HỌ TÊN: HOÀNG ĐẮC PHÚ
NGÀY: 2026-10-25
CA: Ca ngày
THỜI GIAN: 8h-17h (8H)
CÓ ĂN: 1"""

        records = parse_attendance_messages(roster_text)
        assert len(records) == 4

        wb = get_fresh_workbook()
        audit = apply_attendance_records_to_workbook(wb, records)
        assert audit["employees_count"] == 4
        assert len(audit["errors"]) == 0

        ws = wb.active
        # 1. NGUYỄN NGỌC THAO (row 3): 24/10 (col 11) -> HCN=8 (row 3), TCN=3 (row 4), meal=1 (row 9)
        assert ws.cell(row=3, column=11).value == 8
        assert ws.cell(row=4, column=11).value == 3
        assert ws.cell(row=9, column=11).value == 1

        # 2. NGUYỄN ĐỨC CÔNG (row 10): 24/10 -> HCĐ=8 (row 12), meal=0 (row 16)
        assert ws.cell(row=12, column=11).value == 8
        assert ws.cell(row=16, column=11).value == 0

        # 3. TÔ VĂN HƯNG (row 17): 24/10 -> HCĐ=8 (row 19), TCĐ=3 (row 20), meal=1 (row 23)
        assert ws.cell(row=19, column=11).value == 8
        assert ws.cell(row=20, column=11).value == 3
        assert ws.cell(row=23, column=11).value == 1

        # 4. HOÀNG ĐẮC PHÚ (row 24): 25/10 (col 12, Sunday) -> Ngày CN=8 (row 28), meal=1 (row 30)
        assert ws.cell(row=28, column=12).value == 8
        assert ws.cell(row=30, column=12).value == 1

    def test_tier4_03_messy_raw_batch_handling(self):
        """
        Scenario 3: Realistic messy raw chat containing formatting irregularities,
        missing colons, Zalo desktop timestamps, and a corrupted message.
        """
        messy_chat = """[17:28, 24/10/2026] Thao: HỌ TÊN : NGUYỄN NGỌC THAO
NGÀY: 24/10
CA : Ca ngày 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0

1. HỌ TÊN HOÀNG ĐẮC PHÚ
NGÀY 24/10
CA Ca ngày
THỜI GIAN 8h-20h (11H)
CÓ ĂN 1

HỌ TÊN: TÔ VĂN HƯNG
NGÀY: 24/10
CA: Ca đêm
THỜI GIAN: 20h-8h (11H)
CÓ ĂN: có

Bác ơi em bị ốm nghỉ hôm nay nhé"""

        records = parse_attendance_messages(messy_chat)
        assert len(records) >= 4

        # Valid/Warning records
        assert records[0].normalized_name.upper() == "NGUYỄN NGỌC THAO"
        assert records[1].normalized_name.upper() == "HOÀNG ĐẮC PHÚ"
        assert records[2].normalized_name.upper() == "TÔ VĂN HƯNG"
        assert records[2].meal_count == 1

        # Corrupted entry is isolated as ERROR without crashing batch
        last_rec = records[-1]
        assert last_rec.status == "ERROR"
        assert len(last_rec.messages) > 0

    def test_tier4_04_full_end_to_end_pipeline(self):
        """
        Scenario 4: Complete simulated user journey:
        Input Raw Text -> Parse -> UI Review / Inline Edit -> Apply Excel -> Formula Audit.
        """
        # Step 1: User pastes text
        raw_input = """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0"""

        # Step 2: System parses
        records = parse_attendance_messages(raw_input)
        assert len(records) == 1
        rec = records[0]

        # Step 3: User corrects meal count on UI table (e.g. from 0 to 1)
        rec.meal_count = 1

        # Step 4: Confirm and write to Excel
        wb = get_fresh_workbook()
        audit = apply_attendance_records_to_workbook(wb, [rec])
        assert audit["employees_count"] == 1

        # Step 5: Save workbook and inspect file
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        saved_wb = openpyxl.load_workbook(buf, data_only=False)
        ws = saved_wb.active

        # Assert cell K24 has 8, cell K30 has 1 (the edited meal count)
        assert ws.cell(row=24, column=11).value == 8
        assert ws.cell(row=30, column=11).value == 1
        # Formulas preserved
        assert ws.cell(row=24, column=37).value == "=SUM(F24:AJ24)"
        assert ws.cell(row=24, column=44).value == "=SUM(AK24:AP24)"

    def test_tier4_05_stress_month_end_workload(self):
        """
        Scenario 5: High-Density Workload across 15 shifts throughout the month.
        Verifies batch stability and non-corruption of formulas across date columns.
        """
        records = []
        # Generate 15 distinct attendance records
        dates = ["2026-10-24", "2026-10-25", "2026-10-26", "2026-10-27", "2026-10-28"]
        employees = ["NGUYỄN NGỌC THAO", "NGUYỄN ĐỨC CÔNG", "HOÀNG ĐẮC PHÚ"]

        for emp in employees:
            for dt in dates:
                msg = f"HỌ TÊN: {emp}\nNGÀY: {dt}\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nCÓ ĂN: 1"
                records.append(parse_single_message(msg))

        assert len(records) == 15

        wb = get_fresh_workbook()
        audit = apply_attendance_records_to_workbook(wb, records)
        assert audit["employees_count"] == 3
        assert audit["updated_cells"] >= 15
        assert len(audit["errors"]) == 0

        # Verify summary formulas remain intact for all employees
        ws = wb.active
        for r in [3, 10, 24]:
            assert ws.cell(row=r, column=37).value == f"=SUM(F{r}:AJ{r})"
            assert ws.cell(row=r, column=44).value == f"=SUM(AK{r}:AP{r})"
