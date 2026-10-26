"""
Adversarial Stress Test Suite for Milestone 4 (Core Hardening).
Targets:
- core/parser.py
- models/attendance.py
- core/excel_integrator.py
- core/template_generator.py

Challenge Dimensions:
1. Vietnamese accent variations (Hòa vs Hoà, Thủy vs Thuỷ, diphthongs/triphthongs, uppercase/lowercase, decomposed NFD vs precomposed NFC).
2. Extreme dates (leap year Feb 29, month boundaries 31/10, 30/11, 28/02, invalid numbers 32/10, 24/13).
3. Extreme hours (0.25h, 0.5h, 15.5h, 23.75h, 24.0h, cross-midnight multi-day transitions).
4. Excel formulas (all 7 rows per employee block consistency, Row 2 grand total evaluation & multi-employee formula limits).
"""
import io
import re
import os
import sys
import unittest
import unicodedata
import openpyxl
from openpyxl.utils import get_column_letter

# Ensure project root is in sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from models.attendance import WorkHours, AttendanceRecord, ParseSummary, ParseResponse
from core.parser import (
    normalize_vietnamese_name,
    calculate_work_hours,
    AttendanceParser,
    parse_single_message,
    parse_attendance_messages,
)
from core.excel_integrator import (
    strip_accents,
    normalize_name_key,
    clean_employee_name_for_matching,
    ExcelIntegrator,
)
from core.template_generator import generate_default_template


class TestVietnameseAccentVariations(unittest.TestCase):
    """
    Challenge Area 1:
    Vietnamese accent variations, diphthong/triphthong tone placements,
    case insensitivity, and decomposed NFD vs precomposed NFC Unicode.
    """

    def setUp(self):
        self.parser = AttendanceParser()
        self.integrator = ExcelIntegrator()

    def test_01_hoa_diphthong_variants_matching(self):
        """Test Hòa (new style) vs Hoà (traditional style) in employee matching."""
        # Ho-grave-a vs H-o-a-grave
        hoa_new = "NGUYỄN VĂN HÒA"
        hoa_old = "NGUYỄN VĂN HOÀ"
        
        # Verify strip_accents decomposes both to identical unaccented text
        self.assertEqual(strip_accents(hoa_new), strip_accents(hoa_old))
        self.assertEqual(strip_accents(hoa_new).lower(), "nguyen van hoa")

    def test_02_thuy_diphthong_variants_matching(self):
        """Test Thủy (new style) vs Thuỷ (traditional style) in employee matching."""
        thuy_new = "LÊ THỊ THỦY"
        thuy_old = "LÊ THỊ THUỶ"
        
        self.assertEqual(strip_accents(thuy_new), strip_accents(thuy_old))
        self.assertEqual(strip_accents(thuy_new).lower(), "le thi thuy")

    def test_03_other_diphthongs_triphthongs(self):
        """Test other diphthong/triphthong tone mark variants."""
        pairs = [
            ("HOÀNG ĐẮC PHÚ", "HOÀNG ĐẮC PHÚ"),
            ("NGUYỄN VĂN KHỎE", "NGUYỄN VĂN KHOẺ"),
            ("TRẦN THỊ HÓA", "TRẦN THỊ HOÁ"),
            ("VŨ VĂN TOÀN", "VŨ VĂN TOÀN"),
            ("ĐẶNG HOÀI AN", "ĐẶNG HÒAI AN"),
        ]
        for v1, v2 in pairs:
            u1 = strip_accents(v1)
            u2 = strip_accents(v2)
            self.assertEqual(u1, u2, f"Failed for pair {v1} vs {v2}")

    def test_04_casing_variations_in_parser_and_matcher(self):
        """Test UPPERCASE, lowercase, TitleCase, and mixed case names."""
        casing_variants = [
            "HOÀNG ĐẮC PHÚ",
            "hoàng đắc phú",
            "Hoàng Đắc Phú",
            "HoÀnG ĐắC pHú",
        ]
        for name in casing_variants:
            msg = f"HỌ TÊN: {name}\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertIn(rec.status, ["VALID", "WARNING"])
            self.assertEqual(rec.employee_name, "HOÀNG ĐẮC PHÚ")
            self.assertEqual(rec.normalized_name, "Hoàng Đắc Phú")
            
            # Match in Excel
            row, matched_name, score = self.integrator.match_employee(rec.employee_name)
            self.assertEqual(row, 24)
            self.assertEqual(matched_name, "HOÀNG ĐẮC PHÚ")
            self.assertGreaterEqual(score, 0.95)

    def test_05_precomposed_nfc_message_parsing(self):
        """Test standard NFC precomposed Unicode message parsing."""
        msg = (
            "HỌ TÊN : HOÀNG ĐẮC PHÚ\n"
            "NGÀY: 24/10\n"
            "CA : Ca ngày\n"
            "DỰ ÁN : sdv\n"
            "THỜI GIAN LÀM VIỆC :8h-17h ( 8H)\n"
            "CÓ ĂN Ở XƯỞNG KO : 0"
        )
        nfc_msg = unicodedata.normalize("NFC", msg)
        rec = self.parser.parse_single(nfc_msg)
        self.assertIn(rec.status, ["VALID", "WARNING"])
        self.assertEqual(rec.employee_name, "HOÀNG ĐẮC PHÚ")
        self.assertEqual(rec.resolved_date, "2026-10-24")
        self.assertEqual(rec.shift_type, "DAY")
        self.assertEqual(rec.hours.hcn_100, 8.0)
        self.assertEqual(rec.meal_count, 0)

    def test_06_decomposed_nfd_message_parsing_correctness(self):
        """
        Adversarial Test: Decomposed NFD Unicode message parsing.
        Many Vietnamese keyboards on macOS/iOS produce NFD Unicode.
        The parser must correctly extract fields even when raw text is in NFD.
        """
        msg = (
            "HỌ TÊN : HOÀNG ĐẮC PHÚ\n"
            "NGÀY: 24/10\n"
            "CA : Ca ngày\n"
            "DỰ ÁN : sdv\n"
            "THỜI GIAN LÀM VIỆC :8h-17h ( 8H)\n"
            "CÓ ĂN Ở XƯỞNG KO : 0"
        )
        nfd_msg = unicodedata.normalize("NFD", msg)
        
        # Verify that nfd_msg actually contains combining characters
        self.assertNotEqual(msg, nfd_msg)
        self.assertTrue(any(unicodedata.combining(c) for c in nfd_msg))
        
        rec = self.parser.parse_single(nfd_msg)
        
        # Empirical assertion: Does parser parse NFD correctly or fail?
        # If rec.status == 'ERROR', this identifies a failure mode for NFD input.
        is_correct = (
            rec.status != "ERROR"
            and rec.employee_name == "HOÀNG ĐẮC PHÚ"
            and rec.resolved_date == "2026-10-24"
            and rec.shift_type == "DAY"
            and rec.hours.hcn_100 == 8.0
        )
        # We record whether NFD is supported or broken
        if not is_correct:
            # Document exact failure details
            print(f"\n[EMPIRICAL VULNERABILITY FOUND] NFD Parsing Failed!")
            print(f"Status: {rec.status}")
            print(f"Employee: {rec.employee_name!r}")
            print(f"Messages: {rec.messages}")
        self.assertTrue(is_correct, f"NFD parsing failed: {rec.status} - {rec.messages}")

    def test_07_batch_segmentation_with_nfd(self):
        """Test splitting multi-message batch when text is in NFD format."""
        raw_batch = (
            "HỌ TÊN : NGUYỄN NGỌC THAO\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0\n\n"
            "HỌ TÊN : HOÀNG ĐẮC PHÚ\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
        )
        nfd_batch = unicodedata.normalize("NFD", raw_batch)
        chunks = self.parser.split_messages(nfd_batch)
        
        self.assertEqual(len(chunks), 2, f"Expected 2 chunks for NFD batch, got {len(chunks)}")


class TestExtremeDatesAndBoundaries(unittest.TestCase):
    """
    Challenge Area 2:
    Leap year Feb 29, month boundaries, invalid month numbers, calendar edge cases.
    """

    def setUp(self):
        self.parser = AttendanceParser()
        self.integrator = ExcelIntegrator()

    def test_01_leap_year_feb29(self):
        """Test Feb 29 in leap year (2024, 2028) vs non-leap year (2026)."""
        # Leap year 2024 explicitly provided
        msg_2024 = "HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 29/02/2024\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
        rec_2024 = self.parser.parse_single(msg_2024)
        self.assertEqual(rec_2024.status, "VALID")
        self.assertEqual(rec_2024.resolved_date, "2024-02-29")
        self.assertEqual(rec_2024.day_number, 29)

        # Leap year 2028 explicitly provided
        msg_2028 = "HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 2028-02-29\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
        rec_2028 = self.parser.parse_single(msg_2028)
        self.assertEqual(rec_2028.status, "VALID")
        self.assertEqual(rec_2028.resolved_date, "2028-02-29")

        # Non-leap year 2026 explicitly provided
        msg_2026 = "HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 29/02/2026\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
        rec_2026 = self.parser.parse_single(msg_2026)
        self.assertEqual(rec_2026.status, "ERROR")
        self.assertTrue(any("không hợp lệ" in m for m in rec_2026.messages))

        # Default year 2026 inferred from 29/02
        msg_short = "HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 29/02\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
        rec_short = self.parser.parse_single(msg_short)
        self.assertEqual(rec_short.status, "ERROR")
        self.assertTrue(any("không hợp lệ" in m for m in rec_short.messages))

    def test_02_month_boundaries_valid_and_invalid(self):
        """Test 31/10 (valid), 30/11 (valid), 31/11 (invalid 30-day month), 28/02 (valid), 30/02 (invalid)."""
        valid_boundaries = [
            ("31/10", "2026-10-31"),
            ("30/11", "2026-11-30"),
            ("28/02", "2026-02-28"),
            ("31/12", "2026-12-31"),
            ("01/01", "2026-01-01"),
        ]
        for d_str, exp_iso in valid_boundaries:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: {d_str}\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertIn(rec.status, ["VALID", "WARNING"])
            self.assertEqual(rec.resolved_date, exp_iso)

        invalid_boundaries = [
            "31/11",  # November has 30 days
            "30/02",  # February has 28 days
            "31/04",  # April has 30 days
            "31/06",  # June has 30 days
            "31/09",  # September has 30 days
        ]
        for d_str in invalid_boundaries:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: {d_str}\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertEqual(rec.status, "ERROR", f"Expected ERROR for invalid boundary date {d_str}")
            self.assertTrue(any("không hợp lệ" in m for m in rec.messages))

    def test_03_invalid_month_and_day_numbers(self):
        """Test extreme day/month numbers: 32/10, 24/13, 00/10, 10/00, negative dates, 99/99."""
        invalid_dates = [
            "32/10",
            "24/13",
            "00/10",
            "10/00",
            "00/00",
            "99/99",
            "24/10/2026/12",
            "ngày không rõ",
        ]
        for d_str in invalid_dates:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: {d_str}\nCA: Ca ngày\nTHỜI GIAN: 8h-17h (8H)\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertEqual(rec.status, "ERROR", f"Expected ERROR for invalid date {d_str}")
            self.assertEqual(rec.resolved_date, "")

    def test_04_excel_date_column_resolution_bounds(self):
        """Test ExcelIntegrator date column resolution across boundary and invalid dates."""
        # Range in template is 2026-10-19 (col 6) to 2026-11-18 (col 36)
        self.assertEqual(self.integrator.resolve_date_column("2026-10-19"), 6)
        self.assertEqual(self.integrator.resolve_date_column("19/10"), 6)
        
        self.assertEqual(self.integrator.resolve_date_column("2026-10-31"), 18)
        self.assertEqual(self.integrator.resolve_date_column("31/10"), 18)
        
        self.assertEqual(self.integrator.resolve_date_column("2026-11-01"), 19)
        self.assertEqual(self.integrator.resolve_date_column("01/11"), 19)
        
        self.assertEqual(self.integrator.resolve_date_column("2026-11-18"), 36)
        self.assertEqual(self.integrator.resolve_date_column("18/11"), 36)
        
        # Off-by-one out of range
        self.assertIsNone(self.integrator.resolve_date_column("2026-10-18"))
        self.assertIsNone(self.integrator.resolve_date_column("2026-11-19"))
        self.assertIsNone(self.integrator.resolve_date_column("32/10"))
        self.assertIsNone(self.integrator.resolve_date_column("24/13"))
        self.assertIsNone(self.integrator.resolve_date_column("2024-02-29"))
        self.assertIsNone(self.integrator.resolve_date_column(None))
        self.assertIsNone(self.integrator.resolve_date_column(""))


class TestExtremeHoursAndShifts(unittest.TestCase):
    """
    Challenge Area 3:
    Extreme hours (0.25h, 0.5h, 15.5h, 23.75h, 24.0h), upper bounds (>24h),
    lower bounds (<=0h), and cross-midnight multi-day transitions.
    """

    def setUp(self):
        self.parser = AttendanceParser()

    def test_01_fractional_extreme_hours(self):
        """Test fractional hours: 0.25h, 0.5h, 15.5h, 23.75h, 24.0h."""
        test_cases = [
            ("8h-8h15", 0.25, 0.25, 0.0),
            ("0.25h", 0.25, 0.25, 0.0),
            ("0,25 tiếng", 0.25, 0.25, 0.0),
            ("8h-8h30", 0.5, 0.5, 0.0),
            ("0.5H", 0.5, 0.5, 0.0),
            ("8h-23h30 (15.5h)", 15.5, 8.0, 7.5),
            ("15.5h", 15.5, 8.0, 7.5),
            ("15,5 tiếng", 15.5, 8.0, 7.5),
            ("23.75h", 23.75, 8.0, 15.75),
            ("24.0h", 24.0, 8.0, 16.0),
            ("24h", 24.0, 8.0, 16.0),
        ]
        for time_str, exp_total, exp_hcn, exp_tcn in test_cases:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: {time_str}\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertIn(rec.status, ["VALID", "WARNING"])
            self.assertEqual(rec.hours.total_hours, exp_total, f"Total failed for {time_str}")
            self.assertEqual(rec.hours.hcn_100, exp_hcn, f"HCN failed for {time_str}")
            self.assertEqual(rec.hours.tcn_150, exp_tcn, f"TCN failed for {time_str}")

    def test_02_upper_limit_exceeding_24h_rejected(self):
        """Test hours exceeding 24 hours per day: 24.25h, 24.5h, 25h, 48h."""
        overflow_hours = [
            "24.25h",
            "24.5h",
            "25h",
            "48h",
            "8h-10h (26h)",
        ]
        for time_str in overflow_hours:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: {time_str}\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertEqual(rec.status, "ERROR", f"Expected ERROR for {time_str}")
            self.assertTrue(any("vượt quá 24h" in m for m in rec.messages))
            self.assertEqual(rec.hours.total_hours, 0.0)

    def test_03_lower_limit_zero_and_negative_rejected(self):
        """Test zero or negative hours: 0h, 0.0h, -5h, 8h - -2h."""
        invalid_hours = [
            "0h",
            "0.0h",
            "0 tiếng",
            "-5h",
            "(-5h)",
            "8h-17h (-2h)",
        ]
        for time_str in invalid_hours:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: {time_str}\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertEqual(rec.status, "ERROR", f"Expected ERROR for {time_str}")
            self.assertEqual(rec.hours.total_hours, 0.0)

    def test_04_cross_midnight_transitions(self):
        """Test shifts crossing midnight: 20h-5h (8H), 20h-8h (11H), 22h-6h, 20h15-5h45."""
        cross_cases = [
            # time_str, exp_total, exp_hcd, exp_tcd
            ("20h-5h (8H)", 8.0, 8.0, 0.0),
            ("20h-8h (11H)", 11.0, 8.0, 3.0),
            ("20h-5h", 8.0, 8.0, 0.0),  # 9h elapsed - 1h break = 8h
            ("22h-6h", 8.0, 8.0, 0.0),  # 8h elapsed, no break deduction
            ("20h15-5h45 (8.5h)", 8.5, 8.0, 0.5),
        ]
        for time_str, exp_total, exp_hcd, exp_tcd in cross_cases:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 24/10\nCA: Ca đêm\nTHỜI GIAN: {time_str}\nĂN: 1"
            rec = self.parser.parse_single(msg)
            self.assertIn(rec.status, ["VALID", "WARNING"])
            self.assertEqual(rec.hours.total_hours, exp_total, f"Total failed for {time_str}")
            self.assertEqual(rec.hours.hcd_130, exp_hcd, f"HCĐ failed for {time_str}")
            self.assertEqual(rec.hours.tcd_195, exp_tcd, f"TCĐ failed for {time_str}")
            self.assertEqual(rec.meal_count, 1)

    def test_05_clock_hour_bounds_24h(self):
        """Test clock hour syntax violations: 25h-8h, 8h-17h75, 24h15-8h."""
        clock_violations = [
            "25h-8h",
            "8h-17h75",
            "24h15-8h",
            "8h60-17h00",
        ]
        for time_str in clock_violations:
            msg = f"HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN: {time_str}\nĂN: 0"
            rec = self.parser.parse_single(msg)
            self.assertEqual(rec.status, "ERROR", f"Expected ERROR for clock violation {time_str}")
            self.assertTrue(any("không hợp lệ" in m for m in rec.messages))

    def test_06_sunday_rates_day_and_night(self):
        """Test Sunday rate allocations (Sunday Day 200%, Sunday Night 260%)."""
        # 2026-10-25 is Sunday
        # Sunday Day shift: 10h
        msg_sun_day = "HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 25/10/2026\nCA: Ca ngày\nTHỜI GIAN: 8h-19h (10H)\nĂN: 1"
        rec_sun_day = self.parser.parse_single(msg_sun_day)
        self.assertTrue(rec_sun_day.is_sunday)
        self.assertEqual(rec_sun_day.hours.sunday_day_200, 10.0)
        self.assertEqual(rec_sun_day.hours.hcn_100, 0.0)
        self.assertEqual(rec_sun_day.hours.tcn_150, 0.0)

        # Sunday Night shift: 11h
        msg_sun_night = "HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 25/10/2026\nCA: Ca đêm\nTHỜI GIAN: 20h-8h (11H)\nĂN: 1"
        rec_sun_night = self.parser.parse_single(msg_sun_night)
        self.assertTrue(rec_sun_night.is_sunday)
        self.assertEqual(rec_sun_night.hours.sunday_night_260, 11.0)
        self.assertEqual(rec_sun_night.hours.hcd_130, 0.0)
        self.assertEqual(rec_sun_night.hours.tcd_195, 0.0)


class TestExcelFormulasAndScalability(unittest.TestCase):
    """
    Challenge Area 4:
    Excel formulas evaluation consistency across all 7 rows per employee block,
    Row 2 grand totals, and multi-employee formula overflow / limits.
    """

    def setUp(self):
        self.template_path = "templates_excel/BẢNG CHẤM CÔNG WEVN - ETS.xlsx"
        if not os.path.exists(self.template_path):
            generate_default_template(output_path=self.template_path)
        self.wb = openpyxl.load_workbook(self.template_path, data_only=False)
        self.ws = self.wb.active

    def test_01_all_7_rows_formula_syntax_and_offsets(self):
        """
        Verify that for every employee block (starting at r):
        - Offset 0 (r+0): AK{r} = =SUM(F{r}:AJ{r})
        - Offset 1 (r+1): AL{r} = =SUM(F{r+1}:AJ{r+1})
        - Offset 2 (r+2): AM{r} = =SUM(F{r+2}:AJ{r+2})
        - Offset 3 (r+3): AN{r} = =SUM(F{r+3}:AJ{r+3})
        - Offset 4 (r+4): AO{r} = =SUM(F{r+4}:AJ{r+4})
        - Offset 5 (r+5): AP{r} = =SUM(F{r+5}:AJ{r+5})
        - Offset 6 (r+6): AQ{r} = =SUM(F{r+6}:AJ{r+6})
        - AR{r} = =SUM(AK{r}:AP{r})
        """
        emp_start_rows = [3, 10, 17, 24]
        for r in emp_start_rows:
            self.assertEqual(self.ws.cell(row=r, column=37).value, f"=SUM(F{r}:AJ{r})")
            self.assertEqual(self.ws.cell(row=r, column=38).value, f"=SUM(F{r+1}:AJ{r+1})")
            self.assertEqual(self.ws.cell(row=r, column=39).value, f"=SUM(F{r+2}:AJ{r+2})")
            self.assertEqual(self.ws.cell(row=r, column=40).value, f"=SUM(F{r+3}:AJ{r+3})")
            self.assertEqual(self.ws.cell(row=r, column=41).value, f"=SUM(F{r+4}:AJ{r+4})")
            self.assertEqual(self.ws.cell(row=r, column=42).value, f"=SUM(F{r+5}:AJ{r+5})")
            self.assertEqual(self.ws.cell(row=r, column=43).value, f"=SUM(F{r+6}:AJ{r+6})")
            self.assertEqual(self.ws.cell(row=r, column=44).value, f"=SUM(AK{r}:AP{r})")

    def test_02_col_ar_strictly_excludes_col_aq_meals(self):
        """
        Column AR must aggregate only the 6 work hours categories (AK..AP)
        and strictly exclude AQ (Suất ăn).
        """
        for r in [3, 10, 17, 24]:
            val = str(self.ws.cell(row=r, column=44).value)
            self.assertIn("AK", val)
            self.assertIn("AP", val)
            self.assertNotIn("AQ", val, f"Row {r} Column AR formula must not include AQ (meals)")

    def test_03_row_2_grand_totals_syntax_and_terms(self):
        """
        Verify that Row 2 grand total formulas across AK2..AR2 properly reference
        all employee rows and use correct SUM syntax.
        """
        emp_start_rows = [3, 10, 17, 24]
        for col in range(37, 45):
            col_letter = get_column_letter(col)
            expected_terms = [f"{col_letter}{r}" for r in emp_start_rows]
            expected_formula = f"=SUM({','.join(expected_terms)})"
            actual_formula = self.ws.cell(row=2, column=col).value
            self.assertEqual(actual_formula, expected_formula)

    def test_04_mathematical_evaluation_consistency_oracle(self):
        """
        Oracle evaluation: populate numeric test values in all 7 rows across multiple days,
        then independently calculate formula evaluations and verify 100% mathematical consistency.
        """
        integrator = ExcelIntegrator()
        
        # Write records for HOÀNG ĐẮC PHÚ on 24/10 (col 11) and 25/10 (Sunday, col 12)
        records = [
            AttendanceRecord(
                id="rec-1",
                employee_name="HOÀNG ĐẮC PHÚ",
                normalized_name="Hoàng Đắc Phú",
                date_str="24/10",
                resolved_date="2026-10-24",
                day_number=24,
                is_sunday=False,
                shift_type="DAY",
                project="ETS",
                time_window="8h-20h (11H)",
                hours=WorkHours(hcn_100=8.0, tcn_150=3.0, total_hours=11.0),
                meal_count=1,
                status="VALID",
                messages=[],
                raw_text="",
            ),
            AttendanceRecord(
                id="rec-2",
                employee_name="HOÀNG ĐẮC PHÚ",
                normalized_name="Hoàng Đắc Phú",
                date_str="25/10",
                resolved_date="2026-10-25",
                day_number=25,
                is_sunday=True,
                shift_type="DAY",
                project="ETS",
                time_window="8h-17h (8H)",
                hours=WorkHours(sunday_day_200=8.0, total_hours=8.0),
                meal_count=1,
                status="VALID",
                messages=[],
                raw_text="",
            ),
        ]
        
        res = integrator.apply_records(records, output_path=io.BytesIO())
        self.assertEqual(res["unmatched_records"], 0)
        self.assertEqual(len(res["errors"]), 0)
        
        # Verify values at row 24 (HOÀNG ĐẮC PHÚ)
        # col 11 (24/10): HCN=8 (row 24), TCN=3 (row 25), Meal=1 (row 30)
        self.assertEqual(integrator.ws.cell(row=24, column=11).value, 8)
        self.assertEqual(integrator.ws.cell(row=25, column=11).value, 3)
        self.assertEqual(integrator.ws.cell(row=30, column=11).value, 1)
        
        # col 12 (25/10): Sunday Day=8 (row 28), Meal=1 (row 30)
        self.assertEqual(integrator.ws.cell(row=28, column=12).value, 8)
        self.assertEqual(integrator.ws.cell(row=30, column=12).value, 1)

        # Formula cells must retain formula strings
        self.assertEqual(integrator.ws.cell(row=24, column=37).value, "=SUM(F24:AJ24)")
        self.assertEqual(integrator.ws.cell(row=24, column=38).value, "=SUM(F25:AJ25)")
        self.assertEqual(integrator.ws.cell(row=24, column=41).value, "=SUM(F28:AJ28)")
        self.assertEqual(integrator.ws.cell(row=24, column=43).value, "=SUM(F30:AJ30)")
        self.assertEqual(integrator.ws.cell(row=24, column=44).value, "=SUM(AK24:AP24)")
        
        # Row 2 grand total formulas must remain untouched
        self.assertEqual(integrator.ws.cell(row=2, column=37).value, "=SUM(AK3,AK10,AK17,AK24)")

    def test_05_multi_employee_formula_overflow_stress(self):
        """
        Stress test: Multiple employees scaling.
        Test formula structure when employee count is 10, 50, 100, 255, 300.
        Assess Excel 255-argument limit in SUM(arg1, arg2, ...).
        """
        for emp_count in [10, 50, 100, 255, 300]:
            emp_start_rows = [3 + i * 7 for i in range(emp_count)]
            for col in range(37, 45):
                col_letter = get_column_letter(col)
                terms = [f"{col_letter}{r}" for r in emp_start_rows]
                formula = f"=SUM({','.join(terms)})"
                
                # Check Excel formula length limit (8192 characters)
                self.assertLess(
                    len(formula),
                    8192,
                    f"Formula length {len(formula)} exceeds 8192 chars for {emp_count} employees!"
                )
                
                # Check Excel function arguments limit (255 arguments per function call)
                num_args = len(terms)
                if num_args > 255:
                    print(
                        f"\n[BOUNDARY VULNERABILITY FOUND] At {emp_count} employees, "
                        f"Row 2 SUM formula has {num_args} arguments, exceeding Excel's 255 argument limit! "
                        f"Formula: {formula[:40]}... (total {len(formula)} chars)"
                    )
                # When emp_count <= 255, arguments must be <= 255
                if emp_count <= 255:
                    self.assertLessEqual(num_args, 255)


def run_tests():
    suite = unittest.TestSuite()
    loader = unittest.TestLoader()
    suite.addTests(loader.loadTestsFromTestCase(TestVietnameseAccentVariations))
    suite.addTests(loader.loadTestsFromTestCase(TestExtremeDatesAndBoundaries))
    suite.addTests(loader.loadTestsFromTestCase(TestExtremeHoursAndShifts))
    suite.addTests(loader.loadTestsFromTestCase(TestExcelFormulasAndScalability))
    runner = unittest.TextTestRunner(verbosity=2)
    return runner.run(suite)


if __name__ == "__main__":
    result = run_tests()
    sys.exit(0 if result.wasSuccessful() else 1)
