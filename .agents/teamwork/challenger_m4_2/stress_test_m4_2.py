#!/usr/bin/env python3
"""
Tier 5 Adversarial Stress Testing & Code Coverage Hardening Suite (Challenger M4-2).
Targeting:
  1. Concurrency & Resource Leaks:
     - Rapid repeated export requests (sequential + concurrent)
     - openpyxl file handles and OS file descriptor leak verification
     - Mixed concurrent export and streaming download stress
     - Template upload resource cleanup
  2. Table Inline-Edit Edge Cases:
     - 0h / 0 / 0.0 in table cells (frontend parseFloat simulation + backend API coercion)
     - Negative work hours and meal counts (filtering, non-crash, preservation)
     - Decimal hours with comma ('8,5' vs 8.5) (proper 400 error vs valid float write)
     - Special characters in employee name (XSS payloads, SQL injection, emojis, long strings, punctuation)
  3. End-to-End Full Batch Simulation:
     - 56 diverse employee messages (shift types, dates, meals, annotations, malformed messages)
     - POST /api/parse verification
     - Table inline-editing simulation
     - POST /api/export-excel execution
     - GET /api/download/{token} retrieval
     - Deep Excel structure, cell coordinates, formula preservation, fullCalcOnLoad verification
  4. Launcher & Server Hardening:
     - run.py port discovery, usage checking, dependency validation
     - main.py SPA serving, static mounts, CORS, health endpoint
"""

import os
import sys
import gc
import re
import io
import time
import json
import uuid
import shutil
import tempfile
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Dict, Any

import openpyxl
from fastapi.testclient import TestClient

# Ensure project root in sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from main import app
from models.attendance import AttendanceRecord, WorkHours
from core.excel_integrator import ExcelIntegrator, DEFAULT_TEMPLATE_PATH, DEFAULT_OUTPUT_DIR
import run as run_launcher


def get_process_fds() -> int:
    """Returns the current number of open file descriptors in the process."""
    try:
        if os.path.exists("/dev/fd"):
            return len(os.listdir("/dev/fd"))
        elif os.path.exists("/proc/self/fd"):
            return len(os.listdir("/proc/self/fd"))
    except Exception as e:
        print(f"[WARN] Unable to list /dev/fd: {e}")
    return -1


class TestResultsCollector:
    def __init__(self):
        self.passes = 0
        self.failures = 0
        self.records = []

    def check(self, condition: bool, description: str, details: str = ""):
        if condition:
            self.passes += 1
            print(f"  [ PASS ] {description}")
            self.records.append({"status": "PASS", "description": description, "details": details})
        else:
            self.failures += 1
            print(f"  [ FAIL ] {description} -> {details}")
            self.records.append({"status": "FAIL", "description": description, "details": details})


collector = TestResultsCollector()
client = TestClient(app)


# ==============================================================================
# SECTION 1: CONCURRENCY & RESOURCE LEAKS
# ==============================================================================

def test_concurrency_and_resource_leaks():
    print("\n" + "=" * 70)
    print("SECTION 1: CONCURRENCY & RESOURCE LEAK AUDIT")
    print("=" * 70)

    # 1.1 Baseline File Descriptors
    gc.collect()
    baseline_fds = get_process_fds()
    print(f"[*] Baseline open file descriptors: {baseline_fds}")
    collector.check(baseline_fds > 0, "Baseline process file descriptors measurable", f"FDs={baseline_fds}")

    # Standard payload for export
    sample_records = [
        {
            "id": "rec_001",
            "employee_name": "HOÀNG ĐẮC PHÚ",
            "normalized_name": "HOÀNG ĐẮC PHÚ",
            "date_str": "24/10",
            "resolved_date": "2026-10-24",
            "day_number": 24,
            "is_sunday": False,
            "shift_type": "DAY",
            "project": "sdv",
            "time_window": "8h-17h (8H)",
            "hours": {
                "hcn_100": 8.0,
                "tcn_150": 0.0,
                "hcd_130": 0.0,
                "tcd_195": 0.0,
                "sunday_day_200": 0.0,
                "sunday_night_260": 0.0,
                "total_hours": 8.0
            },
            "meal_count": 0,
            "status": "VALID",
            "messages": [],
            "raw_text": "Sample"
        },
        {
            "id": "rec_002",
            "employee_name": "NGUYỄN NGỌC THAO",
            "normalized_name": "NGUYỄN NGỌC THAO",
            "date_str": "24/10",
            "resolved_date": "2026-10-24",
            "day_number": 24,
            "is_sunday": False,
            "shift_type": "DAY",
            "project": "sdv",
            "time_window": "8h-20h (11H)",
            "hours": {
                "hcn_100": 8.0,
                "tcn_150": 3.0,
                "hcd_130": 0.0,
                "tcd_195": 0.0,
                "sunday_day_200": 0.0,
                "sunday_night_260": 0.0,
                "total_hours": 11.0
            },
            "meal_count": 1,
            "status": "VALID",
            "messages": [],
            "raw_text": "Sample"
        }
    ]

    # 1.2 Rapid Sequential Export (50 iterations)
    tokens = []
    status_codes = []
    print("[*] Executing 50 rapid sequential export requests...")
    for i in range(50):
        res = client.post("/api/export-excel", json={
            "records": sample_records,
            "output_filename": f"seq_stress_{i}.xlsx"
        })
        status_codes.append(res.status_code)
        if res.status_code == 200:
            tokens.append(res.json().get("download_token"))

    collector.check(all(sc == 200 for sc in status_codes), "All 50 sequential export requests returned HTTP 200")
    collector.check(len(tokens) == 50, "All 50 sequential export requests generated valid download tokens")

    gc.collect()
    post_seq_fds = get_process_fds()
    print(f"[*] Post-sequential open file descriptors: {post_seq_fds} (Delta: {post_seq_fds - baseline_fds})")
    collector.check(abs(post_seq_fds - baseline_fds) <= 3,
                    "Sequential exports do not leak open file descriptors",
                    f"Baseline: {baseline_fds}, Current: {post_seq_fds}")

    # 1.3 Rapid Concurrent Export (50 threads)
    print("[*] Executing 50 concurrent export requests across 10 worker threads...")
    concurrent_tokens = []
    concurrent_statuses = []

    def _do_export(idx: int):
        c = TestClient(app)
        r = c.post("/api/export-excel", json={
            "records": sample_records,
            "output_filename": f"concurrent_stress_{idx}.xlsx"
        })
        return r.status_code, r.json() if r.status_code == 200 else {}

    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = [executor.submit(_do_export, i) for i in range(50)]
        for f in as_completed(futures):
            sc, data = f.result()
            concurrent_statuses.append(sc)
            if sc == 200 and "download_token" in data:
                concurrent_tokens.append(data["download_token"])

    collector.check(all(sc == 200 for sc in concurrent_statuses),
                    "All 50 concurrent export requests returned HTTP 200",
                    f"Count={len(concurrent_statuses)}")
    collector.check(len(concurrent_tokens) == 50,
                    "All 50 concurrent export requests returned download tokens",
                    f"Count={len(concurrent_tokens)}")

    gc.collect()
    post_concurrent_fds = get_process_fds()
    print(f"[*] Post-concurrent open file descriptors: {post_concurrent_fds} (Delta: {post_concurrent_fds - baseline_fds})")
    collector.check(abs(post_concurrent_fds - baseline_fds) <= 3,
                    "Concurrent exports do not leak open file descriptors",
                    f"Baseline: {baseline_fds}, Current: {post_concurrent_fds}")

    # 1.4 Streaming Downloads Cleanup
    print("[*] Testing file download stream resource cleanup for 25 tokens...")
    dl_statuses = []
    for t in tokens[:25]:
        res_dl = client.get(f"/api/download/{t}")
        dl_statuses.append(res_dl.status_code)
        _ = res_dl.content

    collector.check(all(sc == 200 for sc in dl_statuses), "All 25 downloads returned HTTP 200")
    gc.collect()
    post_download_fds = get_process_fds()
    print(f"[*] Post-download open file descriptors: {post_download_fds} (Delta: {post_download_fds - baseline_fds})")
    collector.check(abs(post_download_fds - baseline_fds) <= 3,
                    "Streaming downloads close file handles cleanly",
                    f"Baseline: {baseline_fds}, Current: {post_download_fds}")

    # 1.5 Mixed Concurrent Export & Download Stress
    print("[*] Executing mixed concurrent export (25) and download (25) stress...")
    mixed_statuses = []

    def _mixed_op(idx: int):
        c = TestClient(app)
        if idx % 2 == 0:
            r = c.post("/api/export-excel", json={
                "records": sample_records,
                "output_filename": f"mixed_stress_{idx}.xlsx"
            })
            return r.status_code
        else:
            t_choice = tokens[idx % len(tokens)]
            r = c.get(f"/api/download/{t_choice}")
            _ = r.content
            return r.status_code

    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = [executor.submit(_mixed_op, i) for i in range(50)]
        for f in as_completed(futures):
            mixed_statuses.append(f.result())

    collector.check(all(sc == 200 for sc in mixed_statuses),
                    "All 50 mixed concurrent export/download operations returned HTTP 200")
    gc.collect()
    post_mixed_fds = get_process_fds()
    print(f"[*] Post-mixed open file descriptors: {post_mixed_fds} (Delta: {post_mixed_fds - baseline_fds})")
    collector.check(abs(post_mixed_fds - baseline_fds) <= 3,
                    "Mixed concurrency does not leak open file descriptors",
                    f"Baseline: {baseline_fds}, Current: {post_mixed_fds}")

    # 1.6 Template Upload Resource Cleanup
    print("[*] Testing template upload file descriptor cleanup...")
    upload_statuses = []
    with open(DEFAULT_TEMPLATE_PATH, "rb") as f:
        file_bytes = f.read()

    for i in range(10):
        res_up = client.post(
            "/api/template/upload",
            files={"file": (f"test_upload_{i}.xlsx", io.BytesIO(file_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        upload_statuses.append(res_up.status_code)

    collector.check(all(sc == 200 for sc in upload_statuses), "All 10 template uploads returned HTTP 200")
    gc.collect()
    post_upload_fds = get_process_fds()
    print(f"[*] Post-upload open file descriptors: {post_upload_fds} (Delta: {post_upload_fds - baseline_fds})")
    collector.check(abs(post_upload_fds - baseline_fds) <= 3,
                    "Template uploads do not leak file descriptors",
                    f"Baseline: {baseline_fds}, Current: {post_upload_fds}")


# ==============================================================================
# SECTION 2: TABLE INLINE-EDIT EDGE CASES
# ==============================================================================

def test_table_inline_edit_edge_cases():
    print("\n" + "=" * 70)
    print("SECTION 2: TABLE INLINE-EDIT EDGE CASES")
    print("=" * 70)

    # 2.1 User Enters "0h" or "0" in Table Cell
    print("[*] Testing Edge Case 2.1: User enters 0h / 0 / 0.0 in hours...")
    # Numeric 0.0 sent to backend
    res_zero_num = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_zero",
            "employee_name": "HOÀNG ĐẮC PHÚ",
            "date_str": "24/10",
            "hcn_100": 0.0,
            "tcn_150": 0.0,
            "hcd_130": 0.0,
            "meal_count": 0
        }]
    })
    collector.check(res_zero_num.status_code == 200, "Backend handles 0.0 hours without error (HTTP 200)")
    zero_summary = res_zero_num.json().get("summary", {})
    # Only meal_count=0 is written, 0.0 hours are skipped to avoid overwriting formulas or blank cells
    collector.check(zero_summary.get("cells_written", 0) == 1,
                    "0.0 hours are not written as positive work to template cells",
                    f"cells_written={zero_summary.get('cells_written')}")

    # String "0h" sent to backend float field
    res_zero_str = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_zero_str",
            "employee_name": "HOÀNG ĐẮC PHÚ",
            "date_str": "24/10",
            "hcn_100": "0h",
            "meal_count": 0
        }]
    })
    collector.check(res_zero_str.status_code == 400,
                    "Backend returns HTTP 400 Bad Request when string '0h' sent for float field without crashing",
                    f"Status={res_zero_str.status_code}, Detail={res_zero_str.json().get('detail')}")

    # 2.2 Negative Values in Table Cell
    print("[*] Testing Edge Case 2.2: User enters negative values (-5, -1)...")
    res_neg = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_neg",
            "employee_name": "HOÀNG ĐẮC PHÚ",
            "date_str": "24/10",
            "hcn_100": -5.0,
            "tcn_150": -2.0,
            "meal_count": -1
        }]
    })
    collector.check(res_neg.status_code == 200, "Backend handles negative values without 500 crash (HTTP 200)")
    neg_summary = res_neg.json().get("summary", {})
    # apply_records checks `if h.hcn_100 > 0:` -> negative hours are ignored
    collector.check(neg_summary.get("cells_written", 0) == 1,
                    "Negative work hours are strictly filtered out (not written to Excel)",
                    f"cells_written={neg_summary.get('cells_written')}")

    # Verify openpyxl output: Row 24 (HCN 100%), Col 11 (24/10) is NOT -5.0
    t_neg = res_neg.json()["download_token"]
    res_neg_dl = client.get(f"/api/download/{t_neg}")
    wb_neg = openpyxl.load_workbook(io.BytesIO(res_neg_dl.content), data_only=False)
    ws_neg = wb_neg.active
    val_k24_neg = ws_neg.cell(row=24, column=11).value
    collector.check(val_k24_neg != -5.0, "Cell value for negative hour is NOT written as negative work in Excel", f"Val={val_k24_neg}")
    wb_neg.close()

    # 2.3 Decimal Hours with Comma ('8,5')
    print("[*] Testing Edge Case 2.3: Decimal hours with comma ('8,5')...")
    # String with comma
    res_comma = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_comma",
            "employee_name": "HOÀNG ĐẮC PHÚ",
            "date_str": "24/10",
            "hcn_100": "8,5",
            "meal_count": 1
        }]
    })
    collector.check(res_comma.status_code == 400,
                    "String '8,5' is safely caught and returns HTTP 400 Bad Request (0 crashes)",
                    f"Status={res_comma.status_code}")

    # Standard float 8.5 sent
    res_dot = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_dot",
            "employee_name": "HOÀNG ĐẮC PHÚ",
            "date_str": "24/10",
            "hcn_100": 8.5,
            "meal_count": 1
        }]
    })
    collector.check(res_dot.status_code == 200, "Float 8.5 is accepted and processed cleanly (HTTP 200)")
    t_dot = res_dot.json()["download_token"]
    res_dot_dl = client.get(f"/api/download/{t_dot}")
    wb_dot = openpyxl.load_workbook(io.BytesIO(res_dot_dl.content), data_only=False)
    ws_dot = wb_dot.active
    # HOÀNG ĐẮC PHÚ is at Row 24 (HCN 100%), Column 11 (24/10)
    val_k24_dot = ws_dot.cell(row=24, column=11).value
    collector.check(val_k24_dot == 8.5, "Cell value 8.5 written as numeric float in Excel (Row 24, Col 11)", f"Val={val_k24_dot}")
    wb_dot.close()

    # 2.4 Special Characters, XSS, and SQLi in Employee Name
    print("[*] Testing Edge Case 2.4: Special characters, XSS, SQLi in employee name...")
    special_names = [
        "<script>alert('XSS')</script>",
        "<img src=x onerror=alert(1)>",
        "Robert'); DROP TABLE Employees;--",
        "' OR '1'='1",
        "👨‍🏭 HOÀNG ĐẮC PHÚ ⚡",
        "HOÀNG ĐẮC PHÚ <KTV> [Xưởng 3]",
        "TÔ VĂN HƯNG -- 0987654321",
        "A" * 5000,   # Extremely long string
        "   ",        # Whitespace only
        "---***###",  # Punctuation only
        "\x00\x01\x02NullBytes"
    ]

    for name in special_names:
        rec_special = {
            "id": f"rec_{uuid.uuid4().hex[:6]}",
            "employee_name": name,
            "date_str": "24/10",
            "hcn_100": 8.0,
            "meal_count": 0
        }
        res_spec = client.post("/api/export-excel", json={"records": [rec_special]})
        collector.check(res_spec.status_code == 200,
                        f"Server handles adversarial name without crash (HTTP 200): {name[:30]!r}",
                        f"Status={res_spec.status_code}")

    # Specific check: Fuzzy match succeeds for role/annotation/emoji variants
    res_matched_emoji = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_emoji",
            "employee_name": "👨‍🏭 HOÀNG ĐẮC PHÚ ⚡",
            "date_str": "24/10",
            "hcn_100": 8.0,
            "meal_count": 0
        }]
    })
    s_emoji = res_matched_emoji.json().get("summary", {})
    collector.check(s_emoji.get("employees_matched") == 1,
                    "Name with emojis '👨‍🏭 HOÀNG ĐẮC PHÚ ⚡' correctly fuzzy-matches template employee",
                    f"Matched={s_emoji.get('employees_matched')}")

    res_matched_role = client.post("/api/export-excel", json={
        "records": [{
            "id": "rec_role",
            "employee_name": "HOÀNG ĐẮC PHÚ <KTV> [Xưởng 3]",
            "date_str": "24/10",
            "hcn_100": 8.0,
            "meal_count": 0
        }]
    })
    s_role = res_matched_role.json().get("summary", {})
    collector.check(s_role.get("employees_matched") == 1,
                    "Name with role annotation 'HOÀNG ĐẮC PHÚ <KTV> [Xưởng 3]' correctly matches template employee",
                    f"Matched={s_role.get('employees_matched')}")


# ==============================================================================
# SECTION 3: END-TO-END FULL BATCH SIMULATION WITH 50+ DIVERSE MESSAGES
# ==============================================================================

def generate_56_diverse_messages() -> str:
    messages = [
        # 1. Standard day shift (Acceptance Criteria 1)
        """HỌ TÊN : HOÀNG ĐẮC PHÚ
NGÀY: 24/10
CA : Ca ngày 
DỰ ÁN : sdv 
THỜI GIAN LÀM VIỆC :8h-17h ( 8H)
CÓ ĂN Ở XƯỞNG KO : 0""",

        # 2. Day shift with overtime
        """HỌ TÊN: NGUYỄN NGỌC THAO
NGÀY: 24/10
CA: Ca ngày
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 8h-20h (11H)
CÓ ĂN Ở XƯỞNG KO: 1""",

        # 3. Sunday day shift (25/10/2026 is Sunday)
        """HỌ TÊN: NGUYỄN ĐỨC CÔNG
NGÀY: 25/10
CA: Ca ngày
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 8h-17h (8H)
CÓ ĂN Ở XƯỞNG KO: 1""",

        # 4. Standard night shift
        """HỌ TÊN: BÙI VĂN TIẾN
NGÀY: 24/10
CA: Ca đêm
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 20h-5h (8H)
CÓ ĂN Ở XƯỞNG KO: 0""",

        # 5. Night shift with overtime
        """HỌ TÊN: TÔ VĂN HƯNG
NGÀY: 24/10
CA: Ca đêm
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 20h-8h (11H)
CÓ ĂN Ở XƯỞNG KO: 1""",

        # 6. Unaccented name
        """HỌ TÊN: HOANG DAC PHU
NGÀY: 20/10
CA: ca ngay
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 8h-17h
CÓ ĂN Ở XƯỞNG KO: com: 1""",

        # 7. Lowercase name, half day morning shift, meal 'k'
        """HỌ TÊN: nguyễn ngọc thao
NGÀY: 21/10
CA: Ca ngày
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 8h-12h (4H)
CÓ ĂN Ở XƯỞNG KO: k""",

        # 8. Name with role annotation, 12-hour continuous shift
        """HỌ TÊN: NGUYỄN ĐỨC CÔNG (KTV)
NGÀY: 22/10
CA: Ca ngày
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 8h-21h (12H)
CÓ ĂN Ở XƯỞNG KO: có""",

        # 9. Name with phone number noise
        """HỌ TÊN: TÔ VĂN HƯNG - 0912345678
NGÀY: 23/10
CA: Ca đêm
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 20h-5h
CÓ ĂN Ở XƯỞNG KO: 1""",

        # 10. Sunday night shift with overtime (25/10/2026 is Sunday)
        """HỌ TÊN: Hoàng Đắc Phú
NGÀY: 25/10
CA: Ca đêm
DỰ ÁN: sdv
THỜI GIAN LÀM VIỆC: 20h-8h (11H)
CÓ ĂN Ở XƯỞNG KO: 1""",

        # 11 - 34. Regular coverage across October and November dates in template
        """HỌ TÊN: Nguyễn Ngọc Thao\nNGÀY: 26/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Nguyễn Đức Công\nNGÀY: 27/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: Có""",
        """HỌ TÊN: Tô Văn Hưng\nNGÀY: 28/10\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Hoàng Đắc Phú\nNGÀY: 29/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-20h (11H)\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Ngọc Thao\nNGÀY: 30/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Đức Công\nNGÀY: 31/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Tô Văn Hưng\nNGÀY: 01/11\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-8h (11H)\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Hoàng Đắc Phú\nNGÀY: 02/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Nguyễn Ngọc Thao\nNGÀY: 03/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-20h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Đức Công\nNGÀY: 04/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Tô Văn Hưng\nNGÀY: 05/11\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Hoàng Đắc Phú\nNGÀY: 06/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Nguyễn Ngọc Thao\nNGÀY: 07/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Đức Công\nNGÀY: 08/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Tô Văn Hưng\nNGÀY: 09/11\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-8h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Hoàng Đắc Phú\nNGÀY: 10/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Nguyễn Ngọc Thao\nNGÀY: 11/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Đức Công\nNGÀY: 12/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Tô Văn Hưng\nNGÀY: 13/11\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Hoàng Đắc Phú\nNGÀY: 14/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-20h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Ngọc Thao\nNGÀY: 15/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Nguyễn Đức Công\nNGÀY: 16/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: Tô Văn Hưng\nNGÀY: 17/11\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: Hoàng Đắc Phú\nNGÀY: 18/11\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",

        # 35 - 44. Unmatched employees (testing tolerance & non-crashing reporting)
        """HỌ TÊN: PHAN VĂN BẢO\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: TRẦN THỊ MAI\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: LÊ VĂN LONG\nNGÀY: 24/10\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: ĐINH QUANG HẢI\nNGÀY: 24/10\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-8h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: VŨ VĂN HÒA\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-20h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: NGÔ VĂN QUÝ\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: ĐẶNG THỊ THẢO\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: DƯƠNG VĂN TÀI\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",
        """HỌ TÊN: PHẠM QUỐC TOÀN\nNGÀY: 24/10\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 1""",
        """HỌ TÊN: MAI VĂN DŨNG\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",

        # 45 - 48. Malformed entries (Testing WARNING / ERROR status classification)
        """HỌ TÊN: HOÀNG ĐẮC PHÚ\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""", # Missing date
        """NGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""",           # Missing name
        """HỌ TÊN: NGUYỄN NGỌC THAO\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: không nhớ rõ\nCÓ ĂN Ở XƯỞNG KO: 1""", # Invalid time
        """HỌ TÊN: TÔ VĂN HƯNG\nNGÀY: 31/02\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h\nCÓ ĂN Ở XƯỞNG KO: 0""", # Invalid date

        # 49 - 56. Special date formats, colons, bracket notes
        """HỌ TÊN: NGUYỄN ĐỨC CÔNG\nNGÀY: 2026-10-24\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h (8H)\nCÓ ĂN Ở XƯỞNG KO: 1""", # ISO date
        """HỌ TÊN: HOÀNG ĐẮC PHÚ\nNGÀY: 24-10-2026\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h-17h (8H)\nCÓ ĂN Ở XƯỞNG KO: 0""", # DD-MM-YYYY
        """HỌ TÊN: TÔ VĂN HƯNG\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 7h-21h (13H)\nCÓ ĂN Ở XƯỞNG KO: 2""", # 13h, meal 2
        """HỌ TÊN: NGUYỄN NGỌC THAO\nNGÀY: 24/10\nCA: Ca ngày\nTHỜI GIAN LÀM VIỆC: 8h30-17h30 (8H)\nCÓ ĂN Ở XƯỞNG KO: 1""", # Minutes
        """HỌ TÊN：NGUYỄN ĐỨC CÔNG\nNGÀY：24/10\nCA：Ca ngày\nTHỜI GIAN LÀM VIỆC：8h-17h\nCÓ ĂN Ở XƯỞNG KO：1""", # Full-width colon
        """HỌ TÊN HOÀNG ĐẮC PHÚ\nNGÀY 24/10\nCA Ca ngày\nTHỜI GIAN LÀM VIỆC 8h-17h (8H)\nCÓ ĂN Ở XƯỞNG KO 0""", # Missing colon
        """HỌ TÊN: TÔ VĂN HƯNG [Tổ 2 - Cơ Điện]\nNGÀY: 24/10\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-5h\nCÓ ĂN Ở XƯỞNG KO: 1""", # Bracket note
        """HỌ TÊN: NGUYỄN NGỌC THAO\nNGÀY: 25/10\nCA: Ca đêm\nTHỜI GIAN LÀM VIỆC: 20h-8h (11H)\nCÓ ĂN Ở XƯỞNG KO: 1""" # Sunday night OT
    ]
    return "\n\n".join(messages)


def test_full_batch_simulation_50_plus():
    print("\n" + "=" * 70)
    print("SECTION 3: END-TO-END FULL BATCH SIMULATION (50+ DIVERSE MESSAGES)")
    print("=" * 70)

    raw_batch = generate_56_diverse_messages()
    header_count = len(re.findall(r'(?:HỌ\s*TÊN|HO\s*TEN)\s*[:：]?', raw_batch, re.IGNORECASE))
    print(f"[*] Prepared full batch with {header_count} messages.")
    collector.check(header_count >= 50, f"Full batch contains >= 50 messages ({header_count} messages)")

    # 3.1 POST /api/parse
    t0 = time.time()
    res_parse = client.post("/api/parse", json={"text": raw_batch})
    t_parse = time.time() - t0
    collector.check(res_parse.status_code == 200, "POST /api/parse returns HTTP 200", f"Elapsed: {t_parse:.3f}s")

    data = res_parse.json()
    records = data.get("records", [])
    summary = data.get("summary", {})
    print(f"[*] Parsed records: {len(records)}, Summary: {summary}")

    collector.check(len(records) >= 50, "Parser segmented >= 50 records", f"Total: {len(records)}")
    collector.check(summary.get("total") == len(records), "Summary total matches records length")

    # In the parser, records with DD/MM date format have an inferred year notice, which is categorized as WARNING (Cần chú ý).
    # All 52 non-corrupt records are validly parsed and eligible for Excel export!
    valid_or_warning = summary.get("valid", 0) + summary.get("warning", 0)
    collector.check(valid_or_warning >= 50,
                    f"Successfully parsed non-fatal records (VALID + WARNING) >= 50",
                    f"Count={valid_or_warning}")
    collector.check(summary.get("error", 0) == 4,
                    "Exactly 4 corrupted records flagged with ERROR status without crashing",
                    f"Error={summary.get('error')}")

    # 3.2 Simulate Table Inline-Editing on Records
    print("[*] Simulating table inline-editing on records before export...")
    rec_phu = next((r for r in records if "HOÀNG ĐẮC PHÚ" in r.get("employee_name", "") and "24/10" in r.get("date_str", "")), None)
    collector.check(rec_phu is not None, "Found record for HOÀNG ĐẮC PHÚ on 24/10")
    if rec_phu:
        collector.check(rec_phu["hours"]["hcn_100"] == 8.0, "Parsed HCN 100% is 8.0")
        collector.check(rec_phu["meal_count"] == 0, "Parsed meal count is 0")

    # Simulate user inline edits in UI:
    # 1. User updates HOÀNG ĐẮC PHÚ's meal count from 0 to 1
    # 2. User modifies NGUYỄN NGỌC THAO on 24/10 overtime from 3.0 to 4.0
    for r in records:
        if "HOÀNG ĐẮC PHÚ" in r.get("employee_name", "") and "24/10" in r.get("date_str", ""):
            r["meal_count"] = 1
        elif "NGUYỄN NGỌC THAO" in r.get("employee_name", "") and "24/10" in r.get("date_str", ""):
            r["hours"]["tcn_150"] = 4.0

    # 3.3 POST /api/export-excel
    t0 = time.time()
    res_export = client.post("/api/export-excel", json={
        "records": records,
        "output_filename": "E2E_STRESS_50_PLUS.xlsx"
    })
    t_export = time.time() - t0
    collector.check(res_export.status_code == 200, "POST /api/export-excel returns HTTP 200", f"Elapsed: {t_export:.3f}s")

    exp_data = res_export.json()
    token = exp_data.get("download_token")
    collector.check(bool(token), "Export response contains valid download token", f"Token={token}")
    exp_summary = exp_data.get("summary", {})
    print(f"[*] Export Summary: {exp_summary}")

    collector.check(exp_summary.get("employees_matched", 0) >= 4, "Matched all 4 template employees", f"Matched={exp_summary.get('employees_matched')}")
    collector.check(exp_summary.get("cells_written", 0) >= 30, "Cells written count >= 30", f"Cells={exp_summary.get('cells_written')}")
    collector.check(exp_summary.get("unmatched_records", 0) >= 10, "Unmatched records reported safely in summary", f"Unmatched={exp_summary.get('unmatched_records')}")

    # 3.4 GET /api/download/{token}
    res_download = client.get(f"/api/download/{token}")
    collector.check(res_download.status_code == 200, "GET /api/download/{token} returns HTTP 200")
    collector.check("spreadsheetml" in res_download.headers.get("content-type", ""),
                    "Response Content-Type is Excel spreadsheetml")

    # 3.5 Deep Workbook Verification
    wb = openpyxl.load_workbook(io.BytesIO(res_download.content), data_only=False)
    ws = wb.active

    # Check HOÀNG ĐẮC PHÚ:
    # Row 24 = HCN 100%, Row 30 = Suất ăn, Column 11 = Date 24/10
    val_phu_hcn = ws.cell(row=24, column=11).value
    collector.check(val_phu_hcn == 8, "HOÀNG ĐẮC PHÚ 24/10 HCN written correctly as 8 (Row 24, Col 11)", f"Val={val_phu_hcn}")
    val_phu_meal = ws.cell(row=30, column=11).value
    collector.check(val_phu_meal == 1, "HOÀNG ĐẮC PHÚ 24/10 Meal written correctly as 1 (Row 30, Col 11; inline edit reflected)", f"Val={val_phu_meal}")

    # Check NGUYỄN NGỌC THAO:
    # Row 3 = HCN 100%, Row 4 = TCN 150%, Column 11 = Date 24/10
    val_thao_tcn = ws.cell(row=4, column=11).value
    collector.check(val_thao_tcn == 4, "NGUYỄN NGỌC THAO 24/10 TCN written correctly as 4 (Row 4, Col 11; inline edit reflected)", f"Val={val_thao_tcn}")

    # Check Sunday shifts: NGUYỄN ĐỨC CÔNG on 25/10 (Row 10 start + offset 4 = Row 14 'Ngày CN 200%')
    # Column 12 = Date 25/10
    val_cong_cn = ws.cell(row=14, column=12).value
    collector.check(val_cong_cn == 8, "NGUYỄN ĐỨC CÔNG 25/10 Sunday day hours written to row offset 4 (Row 14, Col 12)", f"Val={val_cong_cn}")

    # Check Formula Preservation in Columns AK..AR (Columns 37..44)
    # Employee 1 (Row 3):
    f_ak3 = str(ws.cell(row=3, column=37).value)
    collector.check("SUM" in f_ak3.upper(), "Formula in AK3 preserved (=SUM)", f"AK3={f_ak3}")
    f_ar3 = str(ws.cell(row=3, column=44).value)
    collector.check("SUM" in f_ar3.upper(), "Formula in AR3 (TỔNG CÔNG) preserved (=SUM)", f"AR3={f_ar3}")

    # Check Row 2 Grand Total Formulas (Columns AK..AR)
    f_ak2 = str(ws.cell(row=2, column=37).value)
    collector.check("SUM" in f_ak2.upper(), "Grand total formula in row 2 (AK2) preserved (=SUM)", f"AK2={f_ak2}")
    f_ar2 = str(ws.cell(row=2, column=44).value)
    collector.check("SUM" in f_ar2.upper(), "Grand total formula in row 2 (AR2) preserved (=SUM)", f"AR2={f_ar2}")

    # Check fullCalcOnLoad
    collector.check(wb.calculation.fullCalcOnLoad is True, "Workbook calculation.fullCalcOnLoad is True")
    wb.close()


# ==============================================================================
# SECTION 4: RUN.PY AND MAIN.PY HARDENING
# ==============================================================================

def test_run_py_and_main_py_hardening():
    print("\n" + "=" * 70)
    print("SECTION 4: RUN.PY & MAIN.PY HARDENING")
    print("=" * 70)

    # 4.1 run.py port checking and free port discovery
    is_in_use = run_launcher.is_port_in_use(8000)
    collector.check(isinstance(is_in_use, bool), "run.py: is_port_in_use returns boolean")

    free_p = run_launcher.find_free_port(start_port=8100, max_port=8150)
    collector.check(8100 <= free_p <= 8150, "run.py: find_free_port finds open port in range", f"Port={free_p}")

    # 4.2 main.py static routes, fallback, and error handling
    res_root = client.get("/")
    collector.check(res_root.status_code == 200, "main.py: GET / returns HTTP 200 (serves SPA index.html)")
    collector.check("text/html" in res_root.headers.get("content-type", ""), "main.py: GET / serves text/html")

    res_static = client.get("/static/js/app.js")
    collector.check(res_static.status_code == 200, "main.py: Static file mount serves /static/js/app.js")

    res_health = client.get("/api/health")
    collector.check(res_health.status_code == 200 and res_health.json().get("status") == "ok",
                    "main.py: Health check /api/health returns ok")

    # 4.3 API Edge Error Handlers (400, 404)
    res_empty_export = client.post("/api/export-excel", json={"records": []})
    collector.check(res_empty_export.status_code == 400, "api: Export with empty records returns HTTP 400")

    res_fake_dl = client.get("/api/download/nonexistent_token_12345")
    collector.check(res_fake_dl.status_code == 404, "api: Download with invalid token returns HTTP 404")

    res_bad_upload = client.post("/api/template/upload", files={"file": ("malicious.exe", b"MZ...", "application/octet-stream")})
    collector.check(res_bad_upload.status_code == 400, "api: Template upload with invalid extension returns HTTP 400")

    res_empty_parse = client.post("/api/parse", json={"text": ""})
    collector.check(res_empty_parse.status_code == 200 and len(res_empty_parse.json().get("records", [])) == 0,
                    "api: Parse with empty text returns HTTP 200 with 0 records")


# ==============================================================================
# MAIN TEST RUNNER
# ==============================================================================

def main():
    print("=" * 72)
    print("  TIER 5 ADVERSARIAL STRESS TEST SUITE - CHALLENGER M4-2")
    print("=" * 72)

    try:
        test_concurrency_and_resource_leaks()
        test_table_inline_edit_edge_cases()
        test_full_batch_simulation_50_plus()
        test_run_py_and_main_py_hardening()
    except Exception as e:
        import traceback
        traceback.print_exc()
        collector.check(False, "Unexpected exception during test run", str(e))

    print("\n" + "=" * 72)
    print("  FINAL STRESS TEST SUMMARY")
    print("=" * 72)
    print(f"  Total Checks : {collector.passes + collector.failures}")
    print(f"  Passed       : {collector.passes}")
    print(f"  Failed       : {collector.failures}")

    verdict = "APPROVE" if collector.failures == 0 else "REJECT"
    print(f"  VERDICT      : {verdict}")
    print("=" * 72)

    if collector.failures > 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
