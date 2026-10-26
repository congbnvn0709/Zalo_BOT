"""
Template Generator for Excel Attendance Timesheet.
Generates default templates_excel/BẢNG CHẤM CÔNG WEVN - ETS.xlsx from sample_data.csv
with authentic 44-column grid layout, 7-row repeating employee blocks, merged headers,
professional typography and styling, and live Excel formulas (=SUM).
"""
import os
import csv
import re
from typing import Optional, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DEFAULT_CSV_PATH = os.path.join(PROJECT_ROOT, "sample_data.csv")
DEFAULT_TEMPLATE_PATH = os.path.join(PROJECT_ROOT, "templates_excel", "BẢNG CHẤM CÔNG WEVN - ETS.xlsx")


def generate_default_template(
    csv_path: Optional[str] = None,
    output_path: Optional[str] = None
) -> str:
    """
    Builds the authentic master template workbook from sample_data.csv.
    
    Structure:
    - 44 columns: A..AR
    - Row 1: Primary headers (Metadata, 31 Date columns F..AJ, Summary headers AK..AR)
    - Row 2: Subheaders (Day-of-Week DOW F..AJ, Grand Total formulas AK..AR)
    - Rows 3..30: 4 Employee blocks (7 rows each: HCN, TCN, HCĐ, TCĐ, Ngày CN, Đêm CN, Suất ăn)
    - Merged metadata columns A..D per employee block and A1:E2 headers
    - Live Excel formulas injected in AK..AR per employee and row 2 grand totals
    - Freeze panes at F3
    - fullCalcOnLoad set to True
    
    Returns the absolute path to the generated workbook.
    """
    src_csv = os.path.abspath(csv_path or DEFAULT_CSV_PATH)
    dest_xlsx = os.path.abspath(output_path or DEFAULT_TEMPLATE_PATH)

    if not os.path.exists(src_csv):
        raise FileNotFoundError(f"Source sample data CSV not found at: {src_csv}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CHẤM CÔNG WEVN - ETS"
    wb.calculation.fullCalcOnLoad = True

    # 1. Read CSV raw data
    with open(src_csv, mode="r", encoding="utf-8") as f:
        reader = list(csv.reader(f))

    # Populate grid from CSV
    employee_start_rows: List[int] = []
    
    for r_idx, row in enumerate(reader, start=1):
        # Detect employee start rows (Row with STT or Name in col 3 and r_idx >= 3)
        if r_idx >= 3 and len(row) >= 3 and row[2].strip():
            # Every 7 rows starting at row 3
            if (r_idx - 3) % 7 == 0:
                employee_start_rows.append(r_idx)

        for c_idx, val in enumerate(row, start=1):
            if val != "":
                # In columns AK..AR (37..44), skip static totals from CSV; we will inject live formulas
                if r_idx >= 3 and c_idx >= 37:
                    continue
                # Row 2 totals (AK2..AR2) will also be formula-driven
                if r_idx == 2 and c_idx >= 37:
                    continue

                # Coerce numeric values in date cells (cols 6..36) and STT (col 1)
                try:
                    num = float(val)
                    cell_val = int(num) if num.is_integer() else num
                except ValueError:
                    cell_val = val

                ws.cell(row=r_idx, column=c_idx, value=cell_val)

    if not employee_start_rows:
        # Fallback to standard 4 employees if detection did not match
        employee_start_rows = [3, 10, 17, 24]

    # 2. Inject authentic Excel SUM formulas
    # Horizontal Employee row sums in AK..AR at offset 0 (row r_start)
    for r in employee_start_rows:
        ws.cell(row=r, column=37, value=f"=SUM(F{r}:AJ{r})")          # AK: HCN 100%
        ws.cell(row=r, column=38, value=f"=SUM(F{r+1}:AJ{r+1})")      # AL: TCN 150%
        ws.cell(row=r, column=39, value=f"=SUM(F{r+2}:AJ{r+2})")      # AM: HCĐ 130%
        ws.cell(row=r, column=40, value=f"=SUM(F{r+3}:AJ{r+3})")      # AN: TCĐ 195%
        ws.cell(row=r, column=41, value=f"=SUM(F{r+4}:AJ{r+4})")      # AO: Ngày CN 200%
        ws.cell(row=r, column=42, value=f"=SUM(F{r+5}:AJ{r+5})")      # AP: Đêm CN 260%
        ws.cell(row=r, column=43, value=f"=SUM(F{r+6}:AJ{r+6})")      # AQ: Suất ăn
        # AR: TỔNG CÔNG aggregates AK..AP strictly excluding AQ (meals)
        ws.cell(row=r, column=44, value=f"=SUM(AK{r}:AP{r})")

    # Row 2 Grand Totals across all employees
    for col in range(37, 45):
        col_letter = get_column_letter(col)
        terms = [f"{col_letter}{r}" for r in employee_start_rows]
        if len(terms) <= 250:
            formula = f"=SUM({','.join(terms)})"
        else:
            chunks = [terms[i:i + 250] for i in range(0, len(terms), 250)]
            sub_sums = [f"SUM({','.join(c)})" for c in chunks]
            formula = f"=SUM({','.join(sub_sums)})"
        ws.cell(row=2, column=col, value=formula)

    # 3. Merging Structure
    # Top headers A1:E2
    ws.merge_cells("A1:A2")  # STT
    ws.merge_cells("B1:B2")  # MNV
    ws.merge_cells("C1:C2")  # HỌ VÀ TÊN
    ws.merge_cells("D1:D2")  # ĐỊA ĐIỂM LÀM VIỆC
    ws.merge_cells("E1:E2")  # CÔNG

    # Employee block metadata columns A..D
    for r in employee_start_rows:
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 6, end_column=1)  # STT
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 6, end_column=2)  # MNV
        ws.merge_cells(start_row=r, start_column=3, end_row=r + 6, end_column=3)  # HỌ VÀ TÊN
        ws.merge_cells(start_row=r, start_column=4, end_row=r + 6, end_column=4)  # ĐỊA ĐIỂM

    # 4. Styling, Typography, Colors, Borders
    font_header_1 = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_dow = Font(name="Segoe UI", size=9, bold=True, color="1F4E78")
    font_sunday_dow = Font(name="Segoe UI", size=9, bold=True, color="C00000")
    font_totals_row2 = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9, color="000000")
    font_emp_summary = Font(name="Segoe UI", size=9, bold=True, color="000000")

    fill_navy_main = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_navy_dates = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_navy_summary = PatternFill(start_color="162D4A", end_color="162D4A", fill_type="solid")
    fill_dow_weekday = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_dow_sunday = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_grand_total = PatternFill(start_color="203764", end_color="203764", fill_type="solid")

    fill_sunday_col = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
    fill_summary_col = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")

    thin_border_side = Side(style="thin", color="BFBFBF")
    med_border_side = Side(style="medium", color="7F7F7F")

    std_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    max_r = max(employee_start_rows) + 6
    max_c = 44

    # Determine Sunday columns (Row 2 cell == "CN")
    sunday_cols = set()
    for c in range(6, 37):
        if str(ws.cell(row=2, column=c).value).strip().upper() == "CN":
            sunday_cols.add(c)

    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(row=r, column=c)
            # Default border
            cell.border = std_border

            # Row 1 Headers
            if r == 1:
                cell.font = font_header_1
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if c <= 5:
                    cell.fill = fill_navy_main
                elif 6 <= c <= 36:
                    cell.fill = fill_navy_dates
                else:
                    cell.fill = fill_navy_summary

            # Row 2 Subheaders
            elif r == 2:
                if c <= 5:
                    cell.fill = fill_navy_main
                elif 6 <= c <= 36:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if c in sunday_cols:
                        cell.fill = fill_dow_sunday
                        cell.font = font_sunday_dow
                    else:
                        cell.fill = fill_dow_weekday
                        cell.font = font_dow
                else:
                    cell.fill = fill_grand_total
                    cell.font = font_totals_row2
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            # Data rows
            else:
                cell.font = font_data
                # Alignments
                if c == 1 or c == 2 or c == 4:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c == 3:
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                elif c == 5:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif 6 <= c <= 36:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if c in sunday_cols:
                        cell.fill = fill_sunday_col
                elif 37 <= c <= 44:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = font_emp_summary
                    cell.fill = fill_summary_col

            # Bottom boundary for each employee block
            for emp_start in employee_start_rows:
                if r == emp_start + 6:
                    cell.border = Border(
                        left=cell.border.left,
                        right=cell.border.right,
                        top=cell.border.top,
                        bottom=med_border_side
                    )

    # 5. Row Heights & Column Widths
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 20
    for r in range(3, max_r + 1):
        ws.row_dimensions[r].height = 19

    col_widths = {
        1: 5.5,   # A: STT
        2: 7.5,   # B: MNV
        3: 25.0,  # C: HỌ VÀ TÊN
        4: 15.0,  # D: ĐỊA ĐIỂM
        5: 16.0,  # E: CÔNG
    }
    for c in range(6, 37):
        col_widths[c] = 5.0
    for c in range(37, 44):
        col_widths[c] = 11.0
    col_widths[44] = 13.0  # AR: TỔNG CÔNG

    for c, width in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = width

    # 6. Freeze Panes & Calculation Properties
    ws.freeze_panes = "F3"
    wb.calculation.fullCalcOnLoad = True

    # 7. Save workbook
    os.makedirs(os.path.dirname(dest_xlsx), exist_ok=True)
    wb.save(dest_xlsx)
    return dest_xlsx


if __name__ == "__main__":
    out = generate_default_template()
    print(f"Generated default template: {out}")
